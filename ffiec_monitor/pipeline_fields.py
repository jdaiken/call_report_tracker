"""The "Q3 Pipeline Fields" sheet (user's IDC rank model field list, 2026-10-30): one row per
institution per quarter, columns named exactly as the model expects.

`extract` computes, from one Call Report facsimile, every field the Call Report can supply.
Fields the list marks as IDC's own calculations (IDC rank, cost of equity, NOPAT ROE, operating
margin and its volatility, return on earning assets, cost of adjusted debt, leverage spread,
internal growth of equity, balance-sheet cash flow, net loans % insured deposits and available
FHLB) are left blank: the list says to take them from the IDC export. Identifier fields (FRS ID,
cert, name, state, parent, hold type, period) are filled when the sheet is built.

Units follow the list's Format column: percentages as numbers (11.12 = 11.12%), total assets in
$ millions, the new dollar fields in $000 (as reported). Income items are year-to-date,
annualized (Q3 x 4/3) -- an approximation of the list's "1YR" trailing figures. Averages come
from Schedule RC-K (quarterly averages).
"""
from __future__ import annotations

from datetime import date
from typing import Any, Optional

from .ratios import _pick, _ratio_pct, _sum_present, annualization_factor, extract_total_assets

KEY = "_pipeline_fields"     # where the extracted fields live inside a snapshot's ratios JSON
VERSION = 1                  # bump when extract() changes, so the quarter build re-pulls

# (exact column header, format) in the list's order. format: "id", "int", "text", "pct", "musd", "kusd", "num"
FIELDS: list[tuple[str, str]] = [
    ("FRS ID", "id"),
    ("CERTIFICATE NUMBER", "id"),
    ("INSTITUTION NAME", "text"),
    ("STATE", "text"),
    ("PARENT NUMBER", "id"),
    ("HOLD TYPE", "int"),
    ("YEAR PERIOD", "text"),
    ("RANK OF FINANCIAL RATIOS", "int"),
    ("TOTAL ASSETS DOLLARS IN MILLIONS", "musd"),
    ("TIER I CAPITAL AS A % TIER I ASSETS", "pct"),
    ("TIER I & II CAPITAL % RISK-BASED ASSETS", "pct"),
    ("TIER I CAPITAL AS A % OF RISK-BASED ASSETS", "pct"),
    ("LOAN LOSS RESERVE", "pct"),
    ("LOANS 90 DAYS OR MORE PAST DUE", "pct"),
    ("LOANS NONACCRUAL & RESTRUCTURED + REO", "pct"),
    ("BALANCE SHEET CASH FLOW % TIER 1 CAPITAL", "pct"),
    ("NET LOANS & LEASES % INS DEPS & AVL FHLB", "pct"),
    ("INT-BEARING LIAB AS A % OF EARNING ASSETS", "pct"),
    ("STOCKHOLDER RETURN ON EQUITY", "pct"),
    ("ESTIMATE OF COST OF EQUITY", "pct"),
    ("NET OPERATING PROFIT (AFTER-TAX) RETURN ON EQUITY 1YR", "pct"),
    ("CURRENT YIELD ON INVESTMENTS", "pct"),
    ("LOANS AS A % OF EARNING ASSETS", "pct"),
    ("CURRENT YIELD ON LOANS", "pct"),
    ("LOANS NONCURRENT & RESTRUCTURED + REO", "pct"),
    ("% CHG NONCURRENT & RESTRUCTURED + REO", "pct"),
    ("LOANS TO FINANCE COMMERCIAL RE % LOANS", "pct"),
    ("NONINTEREST INCOME 1YR", "pct"),
    ("NONINTEREST EXPENSE 1YR", "pct"),
    ("OPERATING PROFIT MARGIN", "pct"),
    ("RISK ( SD ) IN OP PROFIT MARGIN", "num"),
    ("RETURN ON EARNING ASSETS (AFTER TAX) 1YR", "pct"),
    ("COST OF ADJUSTED DEBT (AFTER TAX) 1YR", "pct"),
    ("LEVERAGE SPREAD", "pct"),
    ("INTERNAL GROWTH OF EQUITY CAPITAL 1YR", "pct"),
    ("Capital Category", "text"),
    ("LEVERAGE MULTIPLIER", "num"),
    ("NET INCOME AS A % OF AVG TANGIBLE ASSETS", "pct"),
    ("Tier 1 capital before AOCI % Tier 1 assets (or AOCI $)", "kusd"),
    ("HTM securities - amortized cost", "kusd"),
    ("HTM securities - fair value", "kusd"),
    ("Tangible common equity", "kusd"),
    ("Uninsured deposits", "kusd"),
    ("Other borrowings (incl. FHLB advances)", "kusd"),
    ("Market value of liquid assets (cash, fed funds, AFS, HTM at FV)", "kusd"),
]

# Left blank on purpose: IDC's own calculations (take from the IDC export).
IDC_ONLY = {
    "RANK OF FINANCIAL RATIOS", "BALANCE SHEET CASH FLOW % TIER 1 CAPITAL", "NET LOANS & LEASES % INS DEPS & AVL FHLB",
    "ESTIMATE OF COST OF EQUITY", "NET OPERATING PROFIT (AFTER-TAX) RETURN ON EQUITY 1YR", "OPERATING PROFIT MARGIN",
    "RISK ( SD ) IN OP PROFIT MARGIN", "RETURN ON EARNING ASSETS (AFTER TAX) 1YR", "COST OF ADJUSTED DEBT (AFTER TAX) 1YR",
    "LEVERAGE SPREAD", "INTERNAL GROWTH OF EQUITY CAPITAL 1YR",
}

# How each computed field is built (Notes sheet).
SOURCES = {
    "TOTAL ASSETS DOLLARS IN MILLIONS": "RC 2170 / 1,000",
    "TIER I CAPITAL AS A % TIER I ASSETS": "RC-R leverage ratio (7204)",
    "TIER I & II CAPITAL % RISK-BASED ASSETS": "RC-R total risk-based ratio (7205); blank for CBLR filers",
    "TIER I CAPITAL AS A % OF RISK-BASED ASSETS": "RC-R Tier 1 risk-based ratio (7206); blank for CBLR filers",
    "LOAN LOSS RESERVE": "Allowance for credit losses on loans (RC 3123) % Tier 1 capital (RC-R 8274)",
    "LOANS 90 DAYS OR MORE PAST DUE": "RC-N 90+ days past due, still accruing (1407) % Tier 1 capital",
    "LOANS NONACCRUAL & RESTRUCTURED + REO": "RC-N nonaccrual (1403) + OREO (RC 2150) % Tier 1 capital; restructured not "
                                             "separately reported since ASU 2022-02",
    "INT-BEARING LIAB AS A % OF EARNING ASSETS": "RC-K average interest-bearing deposits, fed funds purchased/repos and "
                                                 "other borrowed money % RC-K average earning assets",
    "STOCKHOLDER RETURN ON EQUITY": "Net income (RI 4340), annualized, % total equity (RC 3210)",
    "CURRENT YIELD ON INVESTMENTS": "RI securities interest (B488, B489, 4060), annualized, % RC-K average securities",
    "LOANS AS A % OF EARNING ASSETS": "RC-K average loans % RC-K average earning assets",
    "CURRENT YIELD ON LOANS": "RI loan interest and fees (4010), annualized, % RC-K average loans (3360)",
    "LOANS NONCURRENT & RESTRUCTURED + REO": "Nonaccrual + 90+ past due + OREO % loans + OREO (monitor's noncurrent ratio)",
    "% CHG NONCURRENT & RESTRUCTURED + REO": "Change in nonaccrual + 90+ past due + OREO dollars vs. the prior quarter "
                                             "(blank until a prior quarter with this field is stored)",
    "LOANS TO FINANCE COMMERCIAL RE % LOANS": "RC-C construction (F158, F159) + multifamily (1460) + non-owner-occupied "
                                              "nonfarm nonresidential (F161) + CRE not secured by real estate (2746) % total "
                                              "loans (2122)",
    "NONINTEREST INCOME 1YR": "RI 4079, annualized, % RC-K average earning assets",
    "NONINTEREST EXPENSE 1YR": "RI 4093, annualized, % RC-K average earning assets",
    "Capital Category": "Derived from the PCA ratios: WC well, AC adequately, UC under capitalized (CBLR: WC at 9%+)",
    "LEVERAGE MULTIPLIER": "Total assets / total equity",
    "NET INCOME AS A % OF AVG TANGIBLE ASSETS": "Net income, annualized, % (RC-K average assets - goodwill - other "
                                                "intangibles)",
    "Tier 1 capital before AOCI % Tier 1 assets (or AOCI $)": "AOCI, $000 (RC-R B530)",
    "HTM securities - amortized cost": "RC-B total HTM, amortized cost (1754)",
    "HTM securities - fair value": "RC-B total HTM, fair value (1771)",
    "Tangible common equity": "Total equity (3210) - goodwill (3163) - other intangibles (0426) - perpetual preferred (3838)",
    "Uninsured deposits": "RC-O memo 2, estimated uninsured deposits (5597)",
    "Other borrowings (incl. FHLB advances)": "RC other borrowed money (3190)",
    "Market value of liquid assets (cash, fed funds, AFS, HTM at FV)": "Cash and balances due (0081 + 0071) + fed funds "
        "sold (B987) + reverse repos (B989) + AFS at fair value (1773) + HTM at fair value (1771) + equity securities (JA22)",
}

NONCURRENT_RAW = "_noncurrent_raw"   # $000, for next quarter's % change


def _p(c, code: str) -> Optional[float]:
    """An RC item reported as RCFD (consolidated filers) or RCON (everyone else)."""
    return _pick(c, f"RCFD{code}", f"RCON{code}")


def capital_category(lev, t1, tot, cet1) -> Optional[str]:
    if lev is None:
        return None
    if t1 is None and tot is None:          # community bank leverage ratio filer
        return "WC" if lev >= 9 else ("AC" if lev >= 4 else "UC")
    vals = [(lev, 5, 4), (t1, 8, 6), (tot, 10, 8), (cet1, 6.5, 4.5)]
    vals = [v for v in vals if v[0] is not None]
    if all(v >= w for v, w, _a in vals):
        return "WC"
    if all(v >= a for v, _w, a in vals):
        return "AC"
    return "UC"


def extract(facsimile: dict[str, Any], period_end: date) -> dict[str, Any]:
    c = facsimile.get("concepts", facsimile)
    f = annualization_factor(period_end)
    out: dict[str, Any] = {"_v": VERSION}

    assets = extract_total_assets(facsimile)
    out["TOTAL ASSETS DOLLARS IN MILLIONS"] = round(assets / 1000, 3) if assets is not None else None
    lev = _pick(c, "RCOA7204", "RCFA7204")
    tot = _pick(c, "RCOA7205", "RCFA7205")
    t1r = _pick(c, "RCOA7206", "RCFA7206")
    cet1 = _pick(c, "RCOAP793", "RCFAP793")
    out["TIER I CAPITAL AS A % TIER I ASSETS"] = lev
    out["TIER I & II CAPITAL % RISK-BASED ASSETS"] = tot
    out["TIER I CAPITAL AS A % OF RISK-BASED ASSETS"] = t1r

    tier1 = _pick(c, "RCFA8274", "RCOA8274")
    allowance = _p(c, "3123")
    past_due_90 = _p(c, "1407")
    nonaccrual = _p(c, "1403")
    oreo = _p(c, "2150")
    out["LOAN LOSS RESERVE"] = _ratio_pct(allowance, tier1)
    out["LOANS 90 DAYS OR MORE PAST DUE"] = _ratio_pct(past_due_90, tier1)
    out["LOANS NONACCRUAL & RESTRUCTURED + REO"] = _ratio_pct(
        (nonaccrual or 0) + (oreo or 0), tier1) if (nonaccrual is not None or oreo is not None) else None

    avg_loans = _sum_present(c, "RCON3360", "RCFN3360")
    avg_secs = _sum_present(c, "RCONB558", "RCONB559", "RCONB560")
    avg_ea = _sum_present(c, "RCON3360", "RCFN3360", "RCONB558", "RCONB559", "RCONB560", "RCON3365", "RCON3381", "RCON3401")
    avg_ibl = _sum_present(c, "RCON3485", "RCONB563", "RCONHK16", "RCONHK17", "RCFN3404", "RCON3353", "RCON3355")
    out["INT-BEARING LIAB AS A % OF EARNING ASSETS"] = _ratio_pct(avg_ibl, avg_ea)

    net_income = _pick(c, "RIAD4340")
    equity = _p(c, "3210")
    ann = (lambda v: v * f if v is not None else None)            # noqa: E731
    out["STOCKHOLDER RETURN ON EQUITY"] = _ratio_pct(ann(net_income), equity)
    out["CURRENT YIELD ON INVESTMENTS"] = _ratio_pct(ann(_sum_present(c, "RIADB488", "RIADB489", "RIAD4060")), avg_secs)
    out["LOANS AS A % OF EARNING ASSETS"] = _ratio_pct(avg_loans, avg_ea)
    out["CURRENT YIELD ON LOANS"] = _ratio_pct(ann(_pick(c, "RIAD4010")), avg_loans)

    loans = _p(c, "2122")
    noncurrent = None
    if any(v is not None for v in (nonaccrual, past_due_90, oreo)):
        noncurrent = (nonaccrual or 0) + (past_due_90 or 0) + (oreo or 0)
    out[NONCURRENT_RAW] = noncurrent
    out["LOANS NONCURRENT & RESTRUCTURED + REO"] = _ratio_pct(noncurrent, (loans or 0) + (oreo or 0)) if loans else None
    cre = _sum_present(c, "RCONF158", "RCONF159", "RCON1460", "RCONF161", "RCON2746", "RCFD2746")
    out["LOANS TO FINANCE COMMERCIAL RE % LOANS"] = _ratio_pct(cre, loans)
    out["NONINTEREST INCOME 1YR"] = _ratio_pct(ann(_pick(c, "RIAD4079")), avg_ea)
    out["NONINTEREST EXPENSE 1YR"] = _ratio_pct(ann(_pick(c, "RIAD4093")), avg_ea)

    out["Capital Category"] = capital_category(lev, t1r, tot, cet1)
    out["LEVERAGE MULTIPLIER"] = round(assets / equity, 2) if assets and equity else None
    goodwill = _p(c, "3163") or 0.0
    intangibles = _p(c, "0426") or 0.0
    avg_assets = _p(c, "3368")
    out["NET INCOME AS A % OF AVG TANGIBLE ASSETS"] = _ratio_pct(
        ann(net_income), avg_assets - goodwill - intangibles) if avg_assets else None

    out["Tier 1 capital before AOCI % Tier 1 assets (or AOCI $)"] = _pick(c, "RCFAB530", "RCOAB530", "RCFDB530", "RCONB530")
    htm_fv = _p(c, "1771")
    out["HTM securities - amortized cost"] = _p(c, "1754")
    out["HTM securities - fair value"] = htm_fv
    out["Tangible common equity"] = (equity - goodwill - intangibles - (_p(c, "3838") or 0.0)) if equity is not None else None
    out["Uninsured deposits"] = _pick(c, "RCON5597")
    out["Other borrowings (incl. FHLB advances)"] = _p(c, "3190")
    liquid = [_p(c, "0081"), _p(c, "0071"), _pick(c, "RCONB987"), _p(c, "B989"), _p(c, "1773"), htm_fv, _p(c, "JA22")]
    out["Market value of liquid assets (cash, fed funds, AFS, HTM at FV)"] = (
        sum(v for v in liquid if v is not None) if any(v is not None for v in liquid) else None)
    return out


def with_prior(fields: dict[str, Any], prior: Optional[dict[str, Any]]) -> dict[str, Any]:
    """Fill % CHG NONCURRENT from the prior quarter's stored raw noncurrent dollars."""
    cur, old = fields.get(NONCURRENT_RAW), (prior or {}).get(NONCURRENT_RAW)
    fields["% CHG NONCURRENT & RESTRUCTURED + REO"] = (
        (cur - old) / old * 100 if isinstance(cur, (int, float)) and isinstance(old, (int, float)) and old else None)
    return fields


def add_sheet(wb, rows: list[dict], title: str, label: str) -> None:
    """Append the pipeline-fields sheet to an openpyxl workbook: the model's exact column
    headers in row 1 (no title rows, so it loads straight into the model), one row per
    institution; then a notes sheet saying how each field is filled. Grey headers = IDC-only."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head, grey = PatternFill("solid", fgColor="1F3A5F"), PatternFill("solid", fgColor="E3E6EA")
    ws = wb.create_sheet(title[:31])
    for c, (h, _fmt) in enumerate(FIELDS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = grey if h in IDC_ONLY else head
        cell.font = Font(bold=True, color="555555" if h in IDC_ONLY else "FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 60
    fmts = {"pct": "0.00", "musd": "#,##0.0", "kusd": "#,##0", "num": "0.00", "int": "0", "id": "0"}
    for r, row in enumerate(rows, 2):
        for c, (h, fmt) in enumerate(FIELDS, 1):
            v = row.get(h)
            if isinstance(v, float) and fmt in ("pct", "num"):
                v = round(v, 4)
            cell = ws.cell(row=r, column=c, value=v)
            if fmt in fmts:
                cell.number_format = fmts[fmt]
    ws.freeze_panes = "D2"
    for c in range(1, len(FIELDS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions["C"].width = 40
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}{max(len(rows) + 1, 2)}"

    ids = {"FRS ID": "RSSD ID", "CERTIFICATE NUMBER": "FDIC cert", "INSTITUTION NAME": "FFIEC Panel of Reporters",
           "STATE": "FFIEC Panel of Reporters", "PARENT NUMBER": "Holding company RSSD (FDIC BankFind RSSDHCR); 0 if none",
           "HOLD TYPE": "2 = bank with a holding company, 0 = bank without one (Call Report filers only; no Y-9C rows)",
           "YEAR PERIOD": "Quarter of the Call Report"}
    ns = wb.create_sheet(f"{title[:24]} notes")
    ns["A1"] = f"{title}: how each field is filled"
    ns["A1"].font = Font(bold=True, size=14)
    ns["A2"] = (f"{label}. Grey headers are IDC's own calculations: blank here, take them from the IDC export. "
                "Percentages are numbers (11.12 = 11.12%); total assets in $ millions; new dollar fields in $000. "
                "Income items are year-to-date, annualized.")
    for c, h in enumerate(["Field (exact column header)", "How it is filled"], 1):
        cell = ns.cell(row=4, column=c, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
    for i, (h, _fmt) in enumerate(FIELDS, 5):
        how = "Blank: IDC-calculated, take from the IDC export" if h in IDC_ONLY else SOURCES.get(h) or ids.get(h, "")
        ns.cell(row=i, column=1, value=h).font = Font(bold=True)
        ns.cell(row=i, column=2, value=how).alignment = Alignment(wrap_text=True, vertical="top")
    ns.column_dimensions["A"].width = 52
    ns.column_dimensions["B"].width = 110
