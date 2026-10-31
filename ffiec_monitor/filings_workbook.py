"""The Call Report workbook email: an .xlsx of the Call Report filings FFIEC shows as submitted in a window
(Eastern): Summary (counts, IntraFi flag counts, every bank on the enforcement or Treasury Desk
list, with links), Filings (one row per filing, capital category, IntraFi flags, Call Report +
UBPR links), Ratios (capital, earnings, asset quality, liquidity, IntraFi per filing) and Ratio
Guide (how each ratio is calculated, normal vs. concerning ranges, the flag rules). The only
attachment is the workbook -- no PDF, no inline images.

Each filing's facsimile is re-pulled so ratios the daily pipeline doesn't store (ROE,
efficiency, ACL/loans, NPA/assets, Texas ratio, liquidity coverage, unrealized securities
losses, CRE concentration) can be computed; if a pull fails, the stored snapshot ratios are
used and the extra columns are left blank. No DB writes.

The daily run sends it for the filings it processed (pipeline.run_daily); for a past
window, see scripts/send_filings_workbook.py.
"""
from __future__ import annotations

import html
import io
import json
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import db, email_report, ratios, simple_email
from .config import EmailConfig
from .ffiec_client import FFIECClient
from .ubpr import ubpr_url
from .ratios import _pick, _ratio_pct, _sum_present, annualization_factor, extract_total_assets

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
FLAG_FILL = PatternFill("solid", fgColor="FDECEC")
LINK_FONT = Font(color="1D5FBF", underline="single")

# (header, ratios-dict key, category). Values are percents, as extract_ratios returns them.
RATIO_COLUMNS = [
    ("Tier 1 leverage", "tier1_leverage_ratio", "Capital"),
    ("CET1", "cet1_ratio", "Capital"),
    ("Tier 1 risk-based", "tier1_rbc_ratio", "Capital"),
    ("Total risk-based", "total_rbc_ratio", "Capital"),
    ("ROA (annualized)", "roaa", "Earnings"),
    ("ROE (annualized)", "roe", "Earnings"),
    ("Net interest margin", "net_interest_margin", "Earnings"),
    ("Efficiency ratio", "efficiency_ratio", "Earnings"),
    ("NPA / assets", "npa_to_assets", "Asset quality"),
    ("Noncurrent loans / loans", "noncurrent_to_loans", "Asset quality"),
    ("Noncurrent loans + OREO / loans + OREO", "noncurrent_loans_oreo_ratio", "Asset quality"),
    ("Net charge-off rate (annualized)", "charge_off_ratio", "Asset quality"),
    ("ACL / loans", "acl_to_loans", "Asset quality"),
    ("Texas ratio", "texas_ratio", "Asset quality"),
    ("Net loans / deposits", "net_loans_leases_to_deposits", "Liquidity"),
    ("Core deposits / deposits", "core_deposits_to_total_deposits", "Liquidity"),
    ("Brokered + reciprocal / deposits", "brokered_reciprocal_deposits_to_total_deposits", "Liquidity"),
    ("Uninsured deposit share", "uninsured_deposit_share", "Liquidity"),
    ("Liquid assets / uninsured deposits", "liquid_to_uninsured", "Liquidity"),
    ("Unrealized securities losses / Tier 1 capital", "unrealized_loss_to_tier1", "Liquidity"),
    ("CRE / total capital", "cre_to_capital", "Concentration"),
    ("Construction & land / total capital", "construction_to_capital", "Concentration"),
    ("Reciprocal deposits / reciprocal cap", "reciprocal_cap_used", "IntraFi"),
]

# IntraFi flags: (label, test on the ratios dict + capital category). See RATIO_GUIDE for why.
FLAG_RULES = [
    ("Near/over reciprocal cap", lambda r, cat: (r.get("reciprocal_cap_used") or 0) >= 80),
    ("Below well capitalized", lambda r, cat: cat in ("AC", "UC")),
    ("40%+ uninsured", lambda r, cat: (r.get("uninsured_deposit_share") or 0) >= 40),
    ("Liquid assets < uninsured", lambda r, cat: r.get("liquid_to_uninsured") is not None
     and r["liquid_to_uninsured"] < 100),
    ("Securities losses 25%+ of Tier 1", lambda r, cat: (r.get("unrealized_loss_to_tier1") or 0) >= 25),
    ("CRE concentration", lambda r, cat: (r.get("cre_to_capital") or 0) >= 300
     or (r.get("construction_to_capital") or 0) >= 100),
]
CATEGORY_NAMES = {"WC": "Well capitalized", "AC": "Adequately capitalized", "UC": "Undercapitalized"}

# Ratio Guide sheet: (ratio, category, how it's calculated, normal range, concerning range, why it matters).
# Call Report item codes in brackets; RCFD/RCON = consolidated/domestic versions of the same line.
RATIO_GUIDE = [
    ("Tier 1 leverage", "Capital", "Tier 1 capital / average total assets for leverage purposes, as reported on "
     "Schedule RC-R Part I [RCOA7204].", "8% - 12%",
     "Below 5% (below well capitalized); CBLR banks below 9%",
     "Core loss-absorbing capital against all assets. The main test for reciprocal-deposit and brokered-deposit eligibility."),
    ("CET1", "Capital", "Common equity Tier 1 capital / risk-weighted assets [RC-R RCOAP793].", "10% - 15%",
     "Below 6.5% (below well capitalized)", "Highest-quality capital against risk-weighted assets. Not reported by CBLR banks."),
    ("Tier 1 risk-based", "Capital", "Tier 1 capital / risk-weighted assets [RC-R RCOA7206].", "11% - 16%",
     "Below 8% (below well capitalized)", "Not reported by CBLR banks."),
    ("Total risk-based", "Capital", "Total capital (Tier 1 + Tier 2) / risk-weighted assets [RC-R RCOA7205].",
     "12% - 17%", "Below 10% (below well capitalized)", "Not reported by CBLR banks."),
    ("Capital category", "Capital", "Well capitalized: leverage 5%+, Tier 1 RBC 8%+, total RBC 10%+, CET1 6.5%+ "
     "(CBLR banks: leverage 9%+). Adequately capitalized: 4% / 6% / 8% / 4.5%. Below that: undercapitalized. "
     "Shown on the Filings tab.", "Well capitalized",
     "Adequately capitalized or worse",
     "Below well capitalized, reciprocal deposits fall under the special cap (12 CFR 337.6(e)) and brokered "
     "deposits need an FDIC waiver (adequately capitalized) or are prohibited (undercapitalized). An enforcement "
     "order can also make a bank not well capitalized regardless of its ratios."),
    ("ROA (annualized)", "Earnings", "Year-to-date net income [RI 4340], annualized, / average total assets "
     "[RC-K 3368].", "0.8% - 1.4%", "Below 0.5%, or negative for 2+ quarters",
     "Overall profitability. Losses eat capital."),
    ("ROE (annualized)", "Earnings", "Year-to-date net income [RI 4340], annualized, / total equity capital "
     "[RC 3210].", "8% - 14%", "Below 5%, or negative", "Return to shareholders; low ROE limits the ability to raise capital."),
    ("Net interest margin", "Earnings", "Year-to-date net interest income [RI 4074], annualized, / average "
     "earning assets (RC-K averages of loans, securities, fed funds and interest-bearing balances).",
     "3.0% - 4.0%", "Below 2.5%, or falling sharply",
     "Core spread income. Compresses when deposit costs rise faster than asset yields."),
    ("Efficiency ratio", "Earnings", "Noninterest expense [RI 4093] / (net interest income [RI 4074] + "
     "noninterest income [RI 4079]).", "55% - 70%", "Above 80%", "Cost to earn a dollar of revenue; lower is better."),
    ("NPA / assets", "Asset quality", "(Nonaccrual loans [RC-N 1403] + loans 90+ days past due and still accruing "
     "[RC-N 1407] + other real estate owned [RC 2150]) / total assets [RC 2170].", "Below 0.75%", "Above 2%",
     "Problem assets relative to the whole balance sheet."),
    ("Noncurrent loans / loans", "Asset quality", "(Nonaccrual [1403] + 90+ days past due [1407]) / total loans "
     "[RC B528].", "Below 1%", "Above 2%", "Share of the loan book not paying."),
    ("Noncurrent loans + OREO / loans + OREO", "Asset quality", "(Nonaccrual [1403] + 90+ days past due [1407] + "
     "OREO [2150]) / (net loans [B529] + OREO).", "Below 1%", "Above 3%",
     "Includes foreclosed property, so it catches problems already taken back."),
    ("Net charge-off rate (annualized)", "Asset quality", "Year-to-date gross charge-offs less recoveries "
     "(Schedule RI-B Part I), annualized, / average loans [RC-K 3360].", "Below 0.25%", "Above 1%, or rising fast",
     "Losses actually taken. The trend matters more than the level."),
    ("ACL / loans", "Asset quality", "Allowance for credit losses on loans [RC 3123] / total loans [RC B528].",
     "1.0% - 1.5%", "Below 0.8%, or below noncurrent loans",
     "Reserve cushion. Thin reserves with rising noncurrent loans mean future charges against earnings."),
    ("Texas ratio", "Asset quality", "(Nonaccrual [1403] + 90+ days past due [1407] + OREO [2150]) / (tangible "
     "equity (equity [3210] - goodwill [3163] - other intangibles [0426]) + ACL [3123]).", "Below 10%",
     "Above 50% (warning); above 100% (serious)",
     "Problem assets against the capital and reserves available to absorb them. Many failed banks were above 100%."),
    ("Net loans / deposits", "Liquidity", "Loans net of unearned income and ACL [RC B529] / total deposits "
     "[RC 2200, domestic + foreign].", "70% - 90%", "Above 100%",
     "How much of the deposit base is lent out; high means less room to meet outflows."),
    ("Core deposits / deposits", "Liquidity", "(Total deposits - brokered deposits [RC-E 2365] - time deposits "
     "over $250,000 [RC-E J474]) / total deposits.", "Above 80%", "Below 65%",
     "Stable funding. Low core funding means reliance on rate-sensitive money."),
    ("Brokered + reciprocal / deposits", "Liquidity", "(Brokered deposits [RC-E 2365] + reciprocal deposits "
     "[RC-E memo JH83]) / total deposits.", "Below 10%", "Above 25%",
     "Dependence on network and wholesale deposits."),
    ("Uninsured deposit share", "Liquidity", "Estimated uninsured deposits [RC-O memo 2, 5597] / total deposits. "
     "Banks under $1 billion don't report 5597; for them it is estimated as balances in accounts over $250,000 "
     "[RC-O memo F051] less $250,000 per account [F052 x 250].", "Below 30%", "40% or more",
     "Deposits most likely to run in a scare (Silicon Valley Bank was about 90%). The reason banks use IntraFi."),
    ("Liquid assets / uninsured deposits", "Liquidity", "(Cash and balances due [0081 + 0071] + fed funds sold "
     "[B987] + reverse repos [B989] + AFS securities at fair value [1773] + HTM securities at fair value [1771] + "
     "equity securities [JA22]) / uninsured deposits.",
     "Above 150%", "Below 100%",
     "Whether the bank could pay out its uninsured depositors from liquid assets at market value. Pledged "
     "securities are included because most secure public-fund deposits that are themselves counted as uninsured."),
    ("Unrealized securities losses / Tier 1 capital", "Liquidity", "-[(AFS fair value [RC-B 1773] - AFS amortized "
     "cost [1772]) + (HTM fair value [1771] - HTM amortized cost [1754])] / Tier 1 capital [RC-R 8274]. "
     "Positive = net loss.", "Below 10%", "25% or more",
     "Losses that become real if securities must be sold to meet outflows. HTM losses don't hit regulatory capital "
     "until they are sold."),
    ("CRE / total capital", "Concentration", "(Construction and land development [RC-C F158 + F159] + multifamily "
     "[1460] + non-owner-occupied nonfarm nonresidential [F161] + loans to finance CRE not secured by real estate "
     "[2746]) / total risk-based capital [RC-R 3792; Tier 1 + ACL for CBLR banks].", "Below 200%",
     "300% or more (regulatory concentration test, with 50%+ growth over 3 years)",
     "The 2006 interagency CRE guidance trigger for heightened supervisory scrutiny."),
    ("Construction & land / total capital", "Concentration", "Construction and land development loans [RC-C F158 + "
     "F159] / total risk-based capital.", "Below 50%", "100% or more (regulatory concentration test)",
     "The riskiest CRE segment; the other half of the 2006 guidance test."),
    ("Reciprocal deposits / reciprocal cap", "IntraFi", "Reciprocal deposits [RC-E memo JH83] / the lesser of $5 "
     "billion or 20% of total liabilities [RC 2948].", "Below 50%",
     "80% or more (near cap); above 100% (the excess counts as brokered)",
     "Under 12 CFR 337.6(e), reciprocal deposits up to the cap are not brokered. Above it, the excess is treated as "
     "brokered. A bank that falls below well capitalized is limited to its special cap (its average reciprocal "
     "deposits over the previous four quarters)."),
]


def extra_ratios(c: dict[str, Any], period_end) -> dict[str, Optional[float]]:
    """Ratios extract_ratios doesn't compute, from the same facsimile concepts."""
    factor = annualization_factor(period_end)
    net_income = _pick(c, "RIAD4340")
    equity = _pick(c, "RCFD3210", "RCON3210")
    nii = _pick(c, "RIAD4074")
    nonint_income = _pick(c, "RIAD4079")
    nonint_expense = _pick(c, "RIAD4093")
    gross_loans = _pick(c, "RCFDB528", "RCONB528")
    acl = _pick(c, "RCFD3123", "RCON3123")
    nonaccrual = _pick(c, "RCFD1403", "RCON1403") or 0.0
    past_due_90 = _pick(c, "RCFD1407", "RCON1407") or 0.0
    oreo = _pick(c, "RCFD2150", "RCON2150") or 0.0
    goodwill = _pick(c, "RCFD3163", "RCON3163") or 0.0
    intangibles = _pick(c, "RCFD0426", "RCON0426") or 0.0
    total_assets = extract_total_assets(c)
    out: dict[str, Optional[float]] = {"total_assets": total_assets}
    out["roe"] = _ratio_pct(net_income * factor, equity) if net_income is not None else None
    revenue = (nii or 0.0) + (nonint_income or 0.0)
    out["efficiency_ratio"] = _ratio_pct(nonint_expense, revenue) if nonint_expense is not None and revenue else None
    out["acl_to_loans"] = _ratio_pct(acl, gross_loans)
    out["noncurrent_to_loans"] = _ratio_pct(nonaccrual + past_due_90, gross_loans) if gross_loans else None
    out["npa_to_assets"] = _ratio_pct(nonaccrual + past_due_90 + oreo, total_assets) if total_assets else None
    tangible_plus_acl = (equity - goodwill - intangibles + (acl or 0.0)) if equity is not None else None
    out["texas_ratio"] = _ratio_pct(nonaccrual + past_due_90 + oreo, tangible_plus_acl) if tangible_plus_acl else None

    # Liquidity coverage: liquid assets at market value vs. uninsured deposits.
    # 5597 is only reported by $1B+ banks; below that, estimate from accounts over $250,000.
    uninsured = _pick(c, "RCON5597")
    big_balances, big_count = _pick(c, "RCONF051"), _pick(c, "RCONF052")
    if uninsured is None and big_balances is not None and big_count is not None:
        uninsured = max(big_balances - 250 * big_count, 0.0)
        domestic, foreign = _pick(c, "RCON2200"), _pick(c, "RCFN2200") or 0.0
        if domestic:
            out["uninsured_deposit_share"] = _ratio_pct(uninsured, domestic + foreign)
    liquid = [_pick(c, "RCFD0081", "RCON0081"), _pick(c, "RCFD0071", "RCON0071"), _pick(c, "RCONB987"),
              _pick(c, "RCFDB989", "RCONB989"), _pick(c, "RCFD1773", "RCON1773"), _pick(c, "RCFD1771", "RCON1771"),
              _pick(c, "RCFDJA22", "RCONJA22")]
    # Pledged securities are NOT subtracted: most back public-fund deposits, which 5597 already
    # counts as uninsured, so subtracting them would count the same collateral against the bank twice.
    if any(v is not None for v in liquid):
        out["liquid_to_uninsured"] = _ratio_pct(sum(v for v in liquid if v is not None), uninsured) if uninsured else None
    else:
        out["liquid_to_uninsured"] = None

    # Unrealized securities losses (positive = net loss), AFS + HTM, against Tier 1 capital.
    tier1 = _pick(c, "RCFA8274", "RCOA8274")
    afs_fv, afs_cost = _pick(c, "RCFD1773", "RCON1773"), _pick(c, "RCFD1772", "RCON1772")
    htm_fv, htm_cost = _pick(c, "RCFD1771", "RCON1771"), _pick(c, "RCFD1754", "RCON1754")
    gains = [fv - cost for fv, cost in ((afs_fv, afs_cost), (htm_fv, htm_cost)) if fv is not None and cost is not None]
    out["unrealized_loss_to_tier1"] = _ratio_pct(-sum(gains), tier1) if gains and tier1 else None

    # CRE concentration (2006 interagency guidance) against total risk-based capital;
    # CBLR banks don't report total capital, so Tier 1 + ACL stands in.
    total_capital = _pick(c, "RCFA3792", "RCOA3792")
    if total_capital is None and tier1 is not None:
        total_capital = tier1 + (acl or 0.0)
    construction = _sum_present(c, "RCONF158", "RCONF159")
    cre = _sum_present(c, "RCONF158", "RCONF159", "RCON1460", "RCONF161", "RCON2746", "RCFD2746")
    out["cre_to_capital"] = _ratio_pct(cre, total_capital) if cre is not None else None
    out["construction_to_capital"] = _ratio_pct(construction, total_capital) if construction is not None else None
    return out


def capital_category(r: dict[str, Any]) -> Optional[str]:
    """'WC' / 'AC' / 'UC' from the stored capital ratios (pipeline_fields.capital_category)."""
    from .pipeline_fields import capital_category as category
    return category(r.get("tier1_leverage_ratio"), r.get("tier1_rbc_ratio"), r.get("total_rbc_ratio"),
                    r.get(ratios.CET1_KEY))


def intrafi_flags(r: dict[str, Any]) -> list[str]:
    cat = capital_category(r)
    return [label for label, test in FLAG_RULES if test(r, cat)]


def _header(ws, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


def _link(cell, url: str, text: str) -> None:
    cell.value, cell.hyperlink, cell.font = text, url, LINK_FONT


def _pct(v) -> Optional[float]:
    return v / 100 if isinstance(v, (int, float)) else None


def build(records: list[dict], label: str, failed: int, pipeline_rows: Optional[list[dict]] = None,
          pipeline_title: str = "Pipeline Fields") -> bytes:
    """`pipeline_rows` (optional): adds the pipeline-fields sheet + its notes after Ratios
    (ffiec_monitor/pipeline_fields.py) -- the Sunday quarter workbook."""
    wb = Workbook()
    flagged = [r for r in records if r["enforcement"] or r["treasury"]]
    flags = {id(r): intrafi_flags(r["ratios"]) for r in records}
    categories = {id(r): capital_category(r["ratios"]) for r in records}

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Call Report filings {label}"
    ws["A1"].font = Font(bold=True, size=14)
    stats = [
        ("Total filings", len(records)),
        ("New filings", sum(1 for r in records if r["kind"] == "New filing")),
        ("Refilings", sum(1 for r in records if r["kind"] == "Refiling")),
        ("On the enforcement list", sum(1 for r in records if r["enforcement"])),
        ("On the Treasury Desk list", sum(1 for r in records if r["treasury"])),
        ("With any IntraFi flag", sum(1 for r in records if flags[id(r)])),
    ] + [(f"   {label_}", sum(1 for r in records if label_ in flags[id(r)])) for label_, _ in FLAG_RULES]
    for i, (k, v) in enumerate(stats, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=not k.startswith(" "))
        ws.cell(row=i, column=2, value=v)
    row = 3 + len(stats) + 1
    note = ("Enforcement list = active orders in the monitor's enforcement_orders table plus the FDIC "
            "enforcement-decisions export (data/fdic_active_orders.csv), matched by RSSD / FDIC cert. "
            "Treasury Desk list = data/treasury-desk.csv, matched by bank name + state. IntraFi flag rules and "
            "every ratio's calculation and normal / concerning range are on the Ratio Guide tab.")
    if failed:
        note += f" {failed} facsimile pull(s) failed; those rows use stored ratios and leave the extra ratios blank."
    ws.cell(row=row, column=1, value=note).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    ws.row_dimensions[row].height = 62
    row += 2
    ws.cell(row=row, column=1, value="Banks on the enforcement or Treasury Desk list").font = Font(bold=True, size=12)
    _header(ws, row + 1, ["Institution", "City", "State", "Type", "Enforcement", "Treasury Desk", "IntraFi flags",
                          "Call Report", "UBPR"])
    ws.freeze_panes = None
    for i, r in enumerate(flagged, row + 2):
        vals = [r["name"], r["city"], r["state"], r["kind"], "Yes" if r["enforcement"] else "",
                "Yes" if r["treasury"] else "", "; ".join(flags[id(r)])]
        for col, v in enumerate(vals, 1):
            ws.cell(row=i, column=col, value=v)
        _link(ws.cell(row=i, column=8), r["call_report_url"], "Call Report")
        _link(ws.cell(row=i, column=9), r["ubpr_url"], "UBPR")
    if not flagged:
        ws.cell(row=row + 2, column=1, value="None")
    for col, w in zip("ABCDEFGHI", (44, 18, 7, 12, 13, 14, 40, 13, 9)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Filings")
    heads = ["Institution", "City", "State", "RSSD", "FDIC cert", "Period", "Type", "Submitted (ET)",
             "Enforcement", "Treasury Desk", "Other lists", "Capital category", "IntraFi flags", "Call Report", "UBPR"]
    _header(ws, 1, heads)
    for i, r in enumerate(records, 2):
        vals = [r["name"], r["city"], r["state"], r["rssd"], r["cert"], r["period"], r["kind"], r["submitted"],
                "Yes" if r["enforcement"] else "", "Yes" if r["treasury"] else "", ", ".join(r["other_lists"]),
                CATEGORY_NAMES.get(categories[id(r)], ""), "; ".join(flags[id(r)])]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=col, value=v)
            if r["enforcement"] or r["treasury"]:
                cell.fill = FLAG_FILL
        ws.cell(row=i, column=6).number_format = "yyyy-mm-dd"
        if flags[id(r)]:
            ws.cell(row=i, column=13).font = Font(bold=True, color="B00020")
        if categories[id(r)] in ("AC", "UC"):
            ws.cell(row=i, column=12).font = Font(bold=True, color="B00020")
        _link(ws.cell(row=i, column=14), r["call_report_url"], "Call Report")
        _link(ws.cell(row=i, column=15), r["ubpr_url"], "UBPR")
    for col, w in zip("ABCDEFGHIJKLMNO", (44, 18, 7, 10, 10, 11, 11, 17, 12, 13, 18, 20, 40, 12, 8)):
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A1:O{len(records) + 1}"

    ws = wb.create_sheet("Ratios")
    heads = ["Institution", "State", "RSSD", "Period", "Total assets ($000)"] + [h for h, _, _ in RATIO_COLUMNS]
    ws.cell(row=1, column=1, value="")
    for col, (_, _, cat) in enumerate(RATIO_COLUMNS, 6):
        ws.cell(row=1, column=col, value=cat).font = Font(bold=True, color="555555")
    _header(ws, 2, heads)
    ws.row_dimensions[2].height = 45
    for i, r in enumerate(records, 3):
        for col, v in enumerate([r["name"], r["state"], r["rssd"], r["period"], r["ratios"].get("total_assets")], 1):
            ws.cell(row=i, column=col, value=v)
        ws.cell(row=i, column=4).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=5).number_format = "#,##0"
        for col, (_, key, _) in enumerate(RATIO_COLUMNS, 6):
            cell = ws.cell(row=i, column=col, value=_pct(r["ratios"].get(key)))
            cell.number_format = "0.00%"
    ws.column_dimensions["A"].width = 44
    for col in range(2, len(heads) + 1):
        ws.column_dimensions[ws.cell(row=2, column=col).column_letter].width = 13
    ws.auto_filter.ref = f"A2:{ws.cell(row=2, column=len(heads)).column_letter}{len(records) + 2}"

    ws = wb.create_sheet("Ratio Guide")
    ws["A1"] = "Ratio Guide"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Ranges are rules of thumb for U.S. community banks, not regulatory standards, except where a rule is "
                "named. Call Report item codes are in brackets. Year-to-date income items are annualized.")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 32
    _header(ws, 4, ["Ratio", "Category", "How it is calculated", "Normal range", "Concerning range", "Why it matters"])
    for i, row_ in enumerate(RATIO_GUIDE, 5):
        for col, v in enumerate(row_, 1):
            cell = ws.cell(row=i, column=col, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).font = Font(bold=True)
        ws.cell(row=i, column=5).font = Font(color="B00020")
    row = 5 + len(RATIO_GUIDE) + 1
    ws.cell(row=row, column=1, value="IntraFi flags (Filings tab)").font = Font(bold=True, size=12)
    _header(ws, row + 1, ["Flag", "Rule", "", "", "", ""])
    flag_rules = [
        ("Near/over reciprocal cap", "Reciprocal deposits are 80% or more of the reciprocal cap."),
        ("Below well capitalized", "Capital category is adequately capitalized or undercapitalized."),
        ("40%+ uninsured", "Uninsured deposits are 40% or more of total deposits."),
        ("Liquid assets < uninsured", "Liquid assets at market value cover less than 100% of uninsured deposits."),
        ("Securities losses 25%+ of Tier 1", "Net unrealized AFS + HTM securities losses are 25% or more of Tier 1 capital."),
        ("CRE concentration", "CRE is 300% or more of total capital, or construction and land loans are 100% or more."),
        ("Enforcement / Treasury Desk", "Shown in their own columns: an active FDIC / Fed / OCC order, or a bank on the "
                                       "Treasury Desk list."),
    ]
    for i, (name, rule) in enumerate(flag_rules, row + 2):
        ws.cell(row=i, column=1, value=name).font = Font(bold=True)
        ws.cell(row=i, column=2, value=rule).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    ws.freeze_panes = "B5"
    for col, w in zip("ABCDEF", (30, 14, 70, 18, 30, 60)):
        ws.column_dimensions[col].width = w

    if pipeline_rows is not None:
        from . import pipeline_fields
        pipeline_fields.add_sheet(wb, pipeline_rows, pipeline_title, label)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def collect(conn, client: FFIECClient, filings: list[simple_email.Filing]) -> tuple[list[dict], int]:
    """One record per filing, flagged against the lists and with ratios from a fresh
    facsimile pull (stored snapshot ratios if the pull fails). Returns (records, failed)."""
    checker = simple_email.ListChecker(conn)
    records, failed = [], 0
    for f in filings:
        meta = db.get_institution_meta(conn, f.id_rssd) or db.InstitutionMeta(
            id_rssd=f.id_rssd, name=f"RSSD {f.id_rssd}", city="", state="", fdic_cert="",
            total_assets=None, as_of_period=f.period_end)
        try:
            fac = client.retrieve_facsimile(f.id_rssd, f.period_end)
            r = ratios.extract_ratios(fac, f.period_end)
            r.update(extra_ratios(fac.get("concepts", fac), f.period_end))
        except Exception as exc:
            failed += 1
            print(f"facsimile pull failed for {f.id_rssd}: {exc}")
            row = conn.execute("SELECT ratios FROM call_report_snapshots WHERE id_rssd = ? AND period_end = ?",
                               [f.id_rssd, f.period_end]).fetchone()
            r = json.loads(row[0]) if row and row[0] else {}
            r["total_assets"] = meta.total_assets
        lists = checker.lists_for(meta)
        records.append({
            "name": meta.name.strip(), "city": meta.city.strip(), "state": meta.state.strip(),
            "rssd": f.id_rssd, "cert": (meta.fdic_cert or "").strip(), "period": f.period_end,
            "kind": "Refiling" if f.refiling else "New filing",
            "submitted": f"{f.submitted:%Y-%m-%d %I:%M %p}" if f.submitted else "",
            "enforcement": "FDIC order" in lists, "treasury": "Treasury Desk" in lists,
            "other_lists": [x for x in lists if x not in ("FDIC order", "Treasury Desk")],
            "call_report_url": email_report.cdr_report_url(f.id_rssd, f.period_end),
            "ubpr_url": ubpr_url(f.id_rssd, f.period_end),
            "ratios": r,
        })
    records.sort(key=lambda x: (not (x["enforcement"] or x["treasury"]), x["name"].upper()))
    return records, failed


def message(records: list[dict], label: str, failed: int, note: str = "") -> tuple[str, str]:
    """(subject, plain-text body) for the workbook email."""
    enf = sum(1 for x in records if x["enforcement"])
    tsy = sum(1 for x in records if x["treasury"])
    flagged = sum(1 for x in records if intrafi_flags(x["ratios"]))
    subject = (f"Call Report workbook {label}: {len(records)} filings, {enf} enforcement, {tsy} Treasury Desk, "
               f"{flagged} IntraFi flags")
    text = (f"Call Report filings {label}\n\nTotal filings: {len(records)}\nOn the enforcement list: {enf}\n"
            f"On the Treasury Desk list: {tsy}\nWith any IntraFi flag: {flagged}\n\n")
    text += ("The attached workbook has four tabs: Summary, Filings (capital category, IntraFi flags, Call Report + "
             "UBPR links), Ratios (capital, earnings, asset quality, liquidity, IntraFi) and Ratio Guide (how each "
             "ratio is calculated and its normal / concerning range)." if records else "No new Call Report filings.")
    if failed:
        text += f"\n\n{failed} facsimile pull(s) failed; those rows use stored ratios."
    if note:
        text += f"\n\n{note}"
    return subject, text


def send(config: EmailConfig, records: list[dict], label: str, failed: int, note: str = "") -> str:
    """Email the workbook as the only attachment (none when there are no filings). Returns the body."""
    subject, text = message(records, label, failed, note)
    files = []
    if records:
        data = build(records, label, failed)
        fname = f"Call-Report-Filings-{label.replace(' ', '-').replace(',', '')}.xlsx"
        files.append((fname, data, "vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        print(f"Workbook {fname} ({len(data):,} bytes)")
    email_report.send_email(config, subject, text,
                            html_body="<div style='font-family:Arial,sans-serif;font-size:13px;'>"
                                      + html.escape(text).replace("\n", "<br>") + "</div>",
                            attachments=files)
    return text
