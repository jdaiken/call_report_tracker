import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

from ffiec_monitor import db, enforcement, network, network_render, ratios

Q3 = date(2026, 9, 30)


def _add(conn, i, r, assets=200_000, period=Q3):
    db.upsert_snapshot(conn, db.Snapshot(str(i), period, datetime.now(), r, datetime.now()))
    db.upsert_institution_meta(conn, db.InstitutionMeta(str(i), f"Bank {i}", "Springfield", "IL", str(9000 + i), assets, period))


def _r(rec, liab=100_000, unins=30_000, deps=90_000, lev=10.0):
    return {"_reciprocal_raw": rec, "_total_liabilities_raw": liab, "_uninsured_raw": unins, "_total_deposits_raw": deps,
            "tier1_leverage_ratio": lev, "total_rbc_ratio": 15.0}


def test_extraction_of_network_fields():
    fac = {"RCON2948": "1000000", "RCONJH83": "150000", "RCON5597": "300000", "RCON2200": "900000", "RCON2170": "1100000"}
    r = ratios.extract_ratios(fac, Q3)
    assert r["_reciprocal_raw"] == 150000 and r["_total_liabilities_raw"] == 1000000
    assert round(r["reciprocal_cap_used"], 2) == 75.0                      # cap = 20% of $1.0B = $200M
    assert round(r["uninsured_deposit_share"], 2) == 33.33
    big = ratios.extract_ratios({"RCON2948": "40000000", "RCONJH83": "4000000"}, Q3)
    assert round(big["reciprocal_cap_used"], 1) == 80.0                    # $5B cap binds, not 20% ($8B)


def test_network_report(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = db.connect(":memory:")
    _add(conn, 1, _r(19_000))                           # 95% of a $20M cap
    _add(conn, 2, _r(25_000))                           # over the cap
    _add(conn, 3, _r(5_000, lev=4.5))                   # below well capitalized -> special cap
    _add(conn, 4, _r(0, unins=60_000))                  # 67% uninsured, no reciprocal -> prospect
    _add(conn, 5, _r(2_000))
    _add(conn, 6, _r(1_000), period=date(2025, 12, 31))  # stale quarter: excluded
    _add(conn, 7, {"tier1_leverage_ratio": 9.0})        # no network fields and no client: counted missing
    src = enforcement.Sources(actions=[enforcement.Action("FDIC", "Bank 5", "IL", "Consent Order", date(2025, 1, 1),
                                                          "u", "", "9005")], loaded={"FDIC": 1})
    net = network.collect(conn, None, date(2026, 10, 26), enforcement_sources=src,
                          prior_lookup=lambda rssd, p: _r(10_000))
    assert net.missing == 1 and len(net.banks) == 5
    s = net.stats()
    assert (s["users"], s["near_cap"], s["over_cap"], s["special"]) == (4, 1, 1, 1)   # only the order bank
    assert [b.bank.name for b in net.near_cap()] == ["Bank 2", "Bank 1"]
    assert [b.bank.name for b in net.prospects()] == ["Bank 4"]
    assert {b.bank.name for b in net.special()} == {"Bank 5"}
    assert next(b for b in net.banks if b.bank.name == "Bank 3").below_wc
    assert net.near_cap()[1].prior_used == 50.0
    net = network.build(net)
    assert net.summary_by == "rules" and "in this report" in net.summary
    assert {"capital", "exposure", "orders", "special", "cushion"} <= set(net.images)
    html = network_render.render_html(net)
    assert "Reciprocal exposure" in html and "cid:net_exposure" in html and "oncentration" not in html
    pdf = network_render.render_pdf(net)
    assert len(re.findall(rb"/Type /Page[^s]", pdf)) == 7
    wb = load_workbook(io.BytesIO(network_render.build_workbook(net)))
    assert wb.sheetnames == ["Summary", "Banks", "5-quarter trend", "Top 20 holders", "Near cap", "Special cap", "Closest to special cap",
                             "High uninsured"]
    assert str(wb["Summary"]["B8"].value).startswith("=COUNTIF(Banks!")


def test_special_cap_estimate_and_closest(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = db.connect(":memory:")
    _add(conn, 1, {**_r(5_000), "tier1_leverage_ratio": 4.5})                      # below WC, no order: no special cap
    _add(conn, 2, {**_r(8_000), "tier1_leverage_ratio": 8.6, "total_rbc_ratio": None})   # CBLR bank under 9%
    _add(conn, 3, {**_r(8_000), "total_rbc_ratio": 10.4})                           # 0.4 pts above 10%
    _add(conn, 4, _r(9_000))                                                         # under an order dated May 2026
    src = enforcement.Sources(actions=[enforcement.Action("FDIC", "Bank 4", "IL", "Consent Order", date(2026, 5, 6),
                                                          "u", "", "9004")], loaded={"FDIC": 1})
    hist = {date(2026, 3, 31): 4_000, date(2025, 12, 31): 3_000, date(2025, 9, 30): 2_000, date(2025, 6, 30): 1_000}
    net = network.collect(conn, None, date(2026, 10, 26), enforcement_sources=src,
                          prior_lookup=lambda rssd, p: _r(hist.get(p, 500)))
    b1 = next(b for b in net.banks if b.bank.name == "Bank 1")
    assert b1.below_wc and not b1.special_cap and b1.special_cap_est is None
    spec = net.special()[0]
    assert spec.bank.name == "Bank 4" and spec.special_cap_est == 2_500 and spec.over_special == 6_500
    assert "Q2 2025 to Q1 2026" in spec.special_basis and "May 2026" in spec.special_basis
    close = net.closest_to_special()
    assert [b.bank.name for b in close][:3] == ["Bank 1", "Bank 2", "Bank 3"]   # below WC without an order: closest
    assert round(close[1].cushion, 2) == -0.4 and close[1].binding == "CBLR leverage (9%)"


def test_large_banks_left_out(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = db.connect(":memory:")
    _add(conn, 1, _r(19_000))
    _add(conn, 2, _r(4_000_000, liab=300_000_000), assets=60_000_000)             # $60B (Raymond James size): out
    db.upsert_institution_meta(conn, db.InstitutionMeta("3", "BANK OF AMERICA, NATIONAL ASSOCIATION", "Charlotte", "NC",
                                                        "3510", 90_000_000, Q3))  # GSIB name: out
    db.upsert_snapshot(conn, db.Snapshot("3", Q3, datetime.now(), _r(1_000), datetime.now()))
    net = network.collect(conn, None, date(2026, 10, 26))
    assert [b.bank.name for b in net.banks] == ["Bank 1"] and len(net.excluded) == 2
    assert not network.excluded_bank("FIRST STATE BANK", 500_000)


def test_history_cache_and_trend(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = db.connect(":memory:")
    _add(conn, 1, _r(30_000))                                                       # 150% of a $20M cap

    class Client:
        calls = 0

        def retrieve_facsimile(self, rssd, period):
            Client.calls += 1
            return {"RCON2948": "100000", "RCONJH83": "20000"}                       # a year of 100% of cap
    look = network.make_history_lookup(conn, Client())
    net = network.collect(conn, None, date(2026, 10, 26), prior_lookup=look)
    b = net.banks[0]
    assert len(b.history) == 5 and Client.calls == 4 and b.gap_change() == 10_000
    assert [x.bank.name for x in net.trend()] == ["Bank 1"]
    network.collect(conn, None, date(2026, 10, 26), prior_lookup=look)
    assert Client.calls == 4                                                         # cached


def test_size_falls_back_when_meta_assets_missing(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = db.connect(":memory:")
    db.upsert_snapshot(conn, db.Snapshot("1", Q3, datetime.now(), _r(1_000, liab=40_000_000), datetime.now()))
    db.upsert_institution_meta(conn, db.InstitutionMeta("1", "BIG REGIONAL BANK", "St Petersburg", "FL", "1", None, Q3))
    _add(conn, 2, _r(1_000))
    net = network.collect(conn, None, date(2026, 10, 26))
    assert net.excluded == ["BIG REGIONAL BANK"] and [b.bank.name for b in net.banks] == ["Bank 2"]


def test_report_banks_rankings_and_treasury(monkeypatch, tmp_path):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = db.connect(":memory:")
    for i, (lev, nc) in enumerate([(4.0, 1.0), (7.0, 6.0), (12.0, 0.5), (9.0, 2.0)], 1):
        _add(conn, i, {**_r(1_000, lev=lev), "noncurrent_loans_oreo_ratio": nc})
    _add(conn, 5, {**_r(1_000, liab=30_000, lev=3.0), "noncurrent_loans_oreo_ratio": 9.0}, assets=40_000)   # under $50M: not ranked
    _add(conn, 6, {**_r(1_000, lev=2.0)}, period=date(2026, 6, 30))                            # last quarter: not ranked
    f = tmp_path / "t.csv"
    f.write_text("name,aliases,state\nBank 3,,IL\nNo Such Bank,,TX\nBig Bank,,\n")
    _add(conn, 7, _r(1_000), assets=60_000_000)                                          # $60B, on the list: kept
    db.upsert_institution_meta(conn, db.InstitutionMeta("7", "BIG BANK", "Springfield", "IL", "9007", 60_000_000, Q3))
    monkeypatch.setattr(network, "TREASURY_FILE", str(f))
    net = network.collect(conn, None, date(2026, 10, 26), prior_lookup=lambda rssd, p: None)
    assert [b.bank.name for b in net.lowest_capital(2)] == ["Bank 1", "Bank 2"]
    assert net.weakest_assets(1)[0].bank.name == "Bank 2"
    assert {b.bank.name for b in net.treasury_list()} == {"Bank 3", "BIG BANK"} and net.treasury_unmatched == ["No Such Bank"]
    assert {b.bank.name for b in net.focus()} == {"Bank 1", "Bank 2", "Bank 3", "Bank 4", "BIG BANK"}


def test_only_this_quarters_filers(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = db.connect(":memory:")
    _add(conn, 1, _r(1_000))                                         # filed Q3
    _add(conn, 2, _r(2_000), period=date(2026, 6, 30))               # only Q2 so far, under order
    src = enforcement.Sources(actions=[enforcement.Action("FDIC", "Bank 2", "IL", "Consent Order", date(2026, 8, 6),
                                                          "u", "", "9002")], loaded={"FDIC": 1})
    net = network.collect(conn, None, date(2026, 10, 26), enforcement_sources=src, prior_lookup=lambda r, p: None)
    assert [b.bank.name for b in net.banks] == ["Bank 1"] and net.waiting == 1
    assert [b.bank.name for b in net.waiting_orders] == ["Bank 2"] and net.under_order() == []
