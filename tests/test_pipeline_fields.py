import io
from datetime import date

from openpyxl import load_workbook

from ffiec_monitor import pipeline_fields as pf

Q3 = date(2026, 9, 30)
FAC = {
    "RCON2170": "1000000", "RCOA7204": "10.5%", "RCOA7205": "14.0%", "RCOA7206": "12.0%", "RCOAP793": "12.0%",
    "RCOA8274": "100000", "RCON3123": "8000", "RCON1407": "1000", "RCON1403": "3000", "RCON2150": "1000",
    "RIAD4340": "6000", "RCON3210": "110000", "RCON3360": "600000", "RCONB558": "200000", "RCON2122": "620000",
    "RIAD4010": "27000", "RCONF158": "20000", "RCONF161": "100000", "RCON1460": "30000",
    "RCON1754": "50000", "RCON1771": "45000", "RCON5597": "300000", "RCON3190": "40000", "RCON0081": "5000",
    "RCON0071": "20000", "RCON1773": "100000", "RCON3163": "5000", "RCON3368": "990000", "RCON3485": "100000",
}


def test_extract_fills_call_report_fields_and_leaves_idc_blank():
    out = pf.extract(FAC, Q3)
    assert out["TOTAL ASSETS DOLLARS IN MILLIONS"] == 1000.0
    assert out["TIER I CAPITAL AS A % TIER I ASSETS"] == 10.5
    assert out["LOAN LOSS RESERVE"] == 8.0                              # 8,000 / 100,000 Tier 1
    assert out["LOANS NONACCRUAL & RESTRUCTURED + REO"] == 4.0
    assert round(out["STOCKHOLDER RETURN ON EQUITY"], 2) == round(6000 * 4 / 3 / 110000 * 100, 2)
    assert round(out["CURRENT YIELD ON LOANS"], 2) == 6.0               # 27,000 x 4/3 / 600,000
    assert round(out["LOANS TO FINANCE COMMERCIAL RE % LOANS"], 2) == round(150000 / 620000 * 100, 2)
    assert out["Capital Category"] == "WC"
    assert out["HTM securities - fair value"] == 45000
    assert out["Tangible common equity"] == 105000
    assert out["Market value of liquid assets (cash, fed funds, AFS, HTM at FV)"] == 170000
    assert all(h not in out for h in pf.IDC_ONLY)


def test_capital_category_and_prior_change():
    assert pf.capital_category(4.5, 9.0, 11.0, 7.0) == "AC"
    assert pf.capital_category(8.5, None, None, None) == "AC"            # CBLR under 9%
    assert pf.capital_category(3.0, 5.0, 7.0, 4.0) == "UC"
    cur = pf.with_prior({pf.NONCURRENT_RAW: 150.0}, {pf.NONCURRENT_RAW: 100.0})
    assert cur["% CHG NONCURRENT & RESTRUCTURED + REO"] == 50.0


def test_filings_workbook_gets_the_pipeline_sheet_with_exact_headers():
    from ffiec_monitor.filings_workbook import build
    rec = {"name": "Alpha Bank", "city": "Town", "state": "IA", "rssd": "1", "cert": "5", "period": Q3,
           "kind": "New filing", "submitted": "", "enforcement": False, "treasury": False, "other_lists": [],
           "call_report_url": "https://x", "ubpr_url": "https://y", "ratios": {"roe": 7.0}}
    row = {**pf.extract(FAC, Q3), "FRS ID": 1, "INSTITUTION NAME": "Alpha Bank", "YEAR PERIOD": "2026Q3"}
    wb = load_workbook(io.BytesIO(build([rec], "Q3 2026", 0, [row], "Q3 Pipeline Fields")))
    assert wb.sheetnames == ["Summary", "Filings", "Ratios", "Q3 Pipeline Fields", "Q3 Pipeline Fields notes"]
    ws = wb["Q3 Pipeline Fields"]
    assert [c.value for c in ws[1]] == [h for h, _ in pf.FIELDS]
    assert ws["C2"].value == "Alpha Bank" and ws["I2"].value == 1000.0
    assert "ROE (annualized)" in [c.value for c in wb["Ratios"][2]]
