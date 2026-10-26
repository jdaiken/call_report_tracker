"""The designed daily email + PDF (ffiec_monitor/report.py, pdf_report.py, narrative.py)."""
import re
from datetime import date
from unittest import mock

from ffiec_monitor import alerts, db, email_report, narrative, ratios
from ffiec_monitor.watch_list import WatchListChange

RUN = date(2026, 10, 26)
Q3 = date(2026, 9, 30)


def _meta(i, assets, name=None):
    return db.InstitutionMeta(id_rssd=str(100000 + i), name=name or f"Community Bank {i}", city="Springfield",
                              state="IL", fdic_cert=str(9000 + i), total_assets=assets, as_of_period=Q3)


def sample_events(n_refilings=55, n_new=80):
    big = _meta(1, 2_400_000, "First Prairie Bank")
    triggers = alerts.classify_new_period({"tier1_leverage_ratio": 5.8, "total_rbc_ratio": 11.2, "roaa": 0.15},
                                          {"tier1_leverage_ratio": 4.6, "total_rbc_ratio": 11.0, "roaa": -0.20},
                                          [], None, None)
    t1 = [t for t in triggers if t.tier == alerts.TIER1]
    t2 = [t for t in triggers if t.tier == alerts.TIER2] or t1[:1]
    snap = ratios.full_snapshot_rows({"tier1_leverage_ratio": 4.6, "total_rbc_ratio": 11.0, "roaa": -0.2})
    return dict(
        watch_list_changes=[email_report.WatchListChangeEvent(big, WatchListChange(big.id_rssd, "ADDED", "Leverage below 5%"))],
        tier1_events=[email_report.Tier1Event(big, Q3, t1, True, snap)],
        watch_list_snapshot=[email_report.WatchListSnapshotEntry(big, snap)],
        tier2_events=[email_report.Tier2Event(_meta(2, 800_000), t2)],
        restatement_events=[email_report.RestatementEvent(_meta(10 + i, 50_000 + i * 1000), date(2026, 6, 30),
                                                          "Net Interest Margin 3.10% to 3.05%") for i in range(n_refilings)],
        new_filing_events=[email_report.NewFilingEvent(_meta(200 + i, 30_000 + i * 2000), Q3) for i in range(n_new)],
        run_date=RUN,
    )


def test_compose_renders_every_section_with_links(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    rep, html, pdf = email_report.compose(**sample_events())
    assert rep.subject == "FFIEC Call Report Monitor · Mon Oct 26 · 1 Tier 1 alert"
    assert rep.summary_by == "rules" and rep.summary.startswith("Tier 1 alert for First Prairie Bank")
    for title in ("Alerts, refilings and other filings", "Watch List added", "Tier 1", "55 refilings", "GOING ON"):
        assert title in html, title
    # long lists: largest first, capped, with a pointer to the CSV
    assert rep.refilings_total == 55 and len(rep.refilings) == 40
    assert rep.refilings[0].bank.assets_raw >= rep.refilings[-1].bank.assets_raw
    assert "Showing" not in html or "attached spreadsheet" in html
    # every institution links to its Call Report and FDIC profile, in the email and the PDF
    assert "cdr.ffiec.gov/public/ViewFacsimileDirect.aspx" in html
    assert "banks.data.fdic.gov/bankfind-suite/bankfind?cert=9001" in html
    assert pdf[:4] == b"%PDF" and b"bankfind?cert=9001" in pdf and b"ViewFacsimileDirect" in pdf
    assert len(re.findall(rb"/Type /Page[^s]", pdf)) >= 2


def test_quiet_day_summary_and_subject(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    ev = sample_events(3, 2)
    ev.update(watch_list_changes=[], tier1_events=[], watch_list_snapshot=[], tier2_events=[])
    rep, html, pdf = email_report.compose(**ev)
    assert rep.subject.endswith("3 refilings, 2 new filings")
    assert rep.summary.startswith("No alerts today")
    assert "Tier 1 alerts" not in html.split("WHAT")[1]          # empty sections are left out
    assert len(re.findall(rb"/Type /Page[^s]", pdf)) == 5      # the PDF is always five pages


def test_claude_summary_is_grounded(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "key")
    ev = sample_events(3, 2)

    def reply(text):
        r = mock.Mock(status_code=200)
        r.json.return_value = {"content": [{"type": "text", "text": text}]}
        return r

    good = "First Prairie Bank crossed a Tier 1 threshold; its assets are $2.4B."
    with mock.patch.object(narrative.requests, "post", return_value=reply(good)):
        rep, html, _ = email_report.compose(**ev)
    assert rep.summary == good and rep.summary_by == "Claude" and "Written by Claude" not in html

    # a number that is not in the data is rejected (after one retry) and the rule-based text is used
    with mock.patch.object(narrative.requests, "post", return_value=reply("Deposits fell 37.5% overnight.")) as post:
        rep, _, _ = email_report.compose(**ev)
    assert post.call_count == 2 and rep.summary_by == "rules"

    with mock.patch.object(narrative.requests, "post", side_effect=OSError("down")):
        rep, _, _ = email_report.compose(**ev)
    assert rep.summary_by == "rules" and "API call failed" in narrative.STATUS


def test_send_email_attaches_csv_and_pdf():
    from ffiec_monitor.config import EmailConfig
    cfg = EmailConfig.__new__(EmailConfig)
    for k, v in (("address", "a@example.com"), ("app_password", "x"), ("recipient", "a@example.com")):
        object.__setattr__(cfg, k, v)
    with mock.patch.object(email_report.smtplib, "SMTP_SSL") as smtp:
        email_report.send_email(cfg, "s", "body", html_body="<p>x</p>", attachment_filename="c.csv",
                                attachment_content="a,b\n", attachments=[("r.pdf", b"%PDF-1.4", "pdf")])
    msg = smtp.return_value.__enter__.return_value.send_message.call_args[0][0]
    names = [p.get_filename() for p in msg.walk() if p.get_filename()]
    assert names == ["c.csv", "r.pdf"]


def test_new_filings_screen_flags_concerns(monkeypatch):
    from ffiec_monitor import screen
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    healthy = {"tier1_leverage_ratio": 10.5, "total_rbc_ratio": 15.0, "roaa": 1.1, "net_interest_margin": 3.6,
               "net_loans_leases_to_deposits": 82, "core_deposits_to_total_deposits": 85,
               "brokered_reciprocal_deposits_to_total_deposits": 4}
    weak = dict(healthy, tier1_leverage_ratio=5.6, roaa=-0.3, brokered_reciprocal_deposits_to_total_deposits=28)
    assert screen.screen(healthy) == []
    flags = screen.screen(weak)
    assert [f.level for f in flags].count("concern") >= 3 and any(f.key == "combo" for f in flags)
    assert screen.screen(None) == []

    ev = sample_events(0, 0)
    ev.update(watch_list_changes=[], tier1_events=[], watch_list_snapshot=[], tier2_events=[],
              new_filing_events=[email_report.NewFilingEvent(_meta(300, 900_000, "Weak Bank"), Q3, weak),
                                 email_report.NewFilingEvent(_meta(301, 5_000_000, "Strong Bank"), Q3, healthy),
                                 email_report.NewFilingEvent(_meta(302, 40_000, "Unknown Bank"), Q3)])
    rep, html, pdf = email_report.compose(**ev)
    assert rep.counts["new"] == 3 and rep.counts["flagged"] == 1
    assert [r.bank.name for r in rep.flagged] == ["Weak Bank"]
    assert [r.bank.name for r in rep.new_filings] == ["Strong Bank", "Unknown Bank"]
    assert "New filings to review" in html and "near the 5% well-capitalized line" in html
    assert rep.subject.endswith("3 new filings (1 flagged)")
    assert "The screen flagged 1" in rep.summary and "Weak Bank" in rep.summary
    import io
    from pypdf import PdfReader
    assert "New filings to review" in "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    csv_text = email_report.events_csv([], ev["new_filing_events"])
    assert "CONCERN: Leverage ratio 5.60%" in csv_text


def test_at_a_glance_analytics_and_inline_charts(monkeypatch):
    from datetime import datetime
    from ffiec_monitor import analytics
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = db.connect(":memory:")
    for i in range(10):                                     # 10 banks; 4 have filed Q3
        conn.execute("INSERT INTO call_report_snapshots VALUES (?, ?, now(), '{}', now())", [str(i), date(2026, 6, 30)])
    for i in range(4):
        conn.execute("INSERT INTO call_report_snapshots VALUES (?, ?, now(), '{}', now())", [str(i), Q3])
        db.log_alert(conn, str(i), Q3, "TIER3", "NEW_PERIOD_NO_TRIGGER", {})
    db.log_alert(conn, "9", date(2026, 6, 30), "TIER3", "RESTATEMENT_SUB_FLOOR", {"reason": "x"})
    a = analytics.compute(conn, datetime.now().date().replace(year=2026, month=10, day=25))
    assert (a.quarter, a.filed, a.universe, a.deadline, a.days_left) == ("Q3 2026", 4, 10, date(2026, 10, 30), 5)
    assert analytics.latest_quarter_end(date(2026, 1, 5)) == date(2025, 12, 31)

    weak = {"tier1_leverage_ratio": 5.6, "roaa": -0.3}
    ev = sample_events(2, 0)
    ev["new_filing_events"] = [email_report.NewFilingEvent(_meta(300, 900_000), Q3, weak)]
    a.days = [(date(2026, 10, 25), 4, 1)]
    rep, html, pdf = email_report.compose(**ev, analytics=a)
    assert "Q3 2026 FILING SEASON" in html and "4</b> of 10 banks" in html and "5 days left" in html
    assert set(rep.images) == {"activity", "flags"} and all(v[:4] == b"\x89PNG" for v in rep.images.values())
    assert 'src="cid:activity"' in html and 'src="cid:flags"' in html
    assert dict((l, r) for l, r, _ in rep.flag_counts)["Capital near the line"] == 1
    assert b"/Subtype /Image" in pdf



def test_zero_alert_tiles_are_hidden(monkeypatch):
    from ffiec_monitor.report import tile_specs
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    quiet = {"tier1": 0, "tier2": 0, "watch": 0, "refilings": 3, "new": 5, "flagged": 1}
    assert [t[0] for t in tile_specs(quiet)] == ["NEW FILINGS", "REFILINGS", "NEW, FLAGGED"]
    busy = dict(quiet, tier1=2, watch=1)
    assert [t[0] for t in tile_specs(busy)][:2] == ["TIER 1 ALERTS", "WATCH LIST"]
    ev = sample_events(3, 2)
    ev.update(watch_list_changes=[], tier1_events=[], watch_list_snapshot=[], tier2_events=[])
    _, html, _ = email_report.compose(**ev)
    assert "TIER 1 ALERTS" not in html and ">TIER 2<" not in html and ">WATCH LIST<" not in html


def test_capital_status_and_enforcement_matching(monkeypatch):
    from ffiec_monitor import enforcement
    from ffiec_monitor.report import capital_status
    assert capital_status({"tier1_leverage_ratio": 9.5, "total_rbc_ratio": 14.0}) == []
    assert capital_status({"tier1_leverage_ratio": 4.8, "total_rbc_ratio": 9.1})[0].startswith("Leverage ratio 4.80%")
    from ffiec_monitor.report import cblr_status
    assert capital_status({"tier1_leverage_ratio": 8.4}) == []                          # CBLR filer: not LTWC
    assert "grace period" in cblr_status({"tier1_leverage_ratio": 8.4})
    assert capital_status(None) == []

    fed = ("Banking Organization,Entity,Individual,Individual Affiliation,Effective Date,Termination Date,Action,Note,URL\n"
           '"First Prairie Bank, Springfield, Illinois",,,,03/01/2025,,Written Agreement,,https://fed.example/1\n'
           '"Old Bank, Dover, Delaware",,,,01/01/2019,06/01/2021,Cease and Desist,,https://fed.example/2\n'
           '"Some Bank, Austin, Texas",,Jane Doe,Some Bank,01/01/2024,,Prohibition,,https://fed.example/3\n')
    acts = enforcement.parse_fed(fed)
    assert [(a.name, a.state, a.action) for a in acts] == [("First Prairie Bank", "IL", "Written Agreement")]
    occ = enforcement.parse_occ([
        {"Institution": "Community Bank 2, National Association", "Location": "Springfield, IL", "Individual": "",
         "StartDate": "2026-02-01T00:00:00", "TerminationDate": None, "TypeDescription": "Formal Agreement"},
        {"Institution": "Gone Bank", "Location": "Austin, TX", "Individual": "", "StartDate": "2020-01-01T00:00:00",
         "TerminationDate": "2022-01-01T00:00:00", "TypeDescription": "Formal Agreement"}])
    assert [(a.name, a.state) for a in occ] == [("Community Bank 2, National Association", "IL")]
    html = ("<table><tr><th>Enforcement Action</th><th>Date</th><th>Matter</th><th>Institution</th></tr>"
            "<tr><td><a href='/x'>In the Matter of Community Bank 10</a></td><td>05/02/2025</td><td>2025-01</td>"
            "<td>Depository Institution</td></tr>"
            "<tr><td>In the Matter of Coin Shop LLC</td><td>05/02/2025</td><td>2025-02</td><td>Money Services Business</td></tr></table>")
    fin = enforcement.parse_fincen(html, date(2026, 10, 26))
    assert [(a.name, a.url) for a in fin] == [("Community Bank 10", "https://www.fincen.gov/x")]

    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    ev = sample_events(3, 0)
    ev["new_filing_events"] = [email_report.NewFilingEvent(_meta(1, 2_400_000, "First Prairie Bank"), Q3,
                                                            {"tier1_leverage_ratio": 4.7, "total_rbc_ratio": 9.0})]
    src = enforcement.Sources(actions=acts + occ + fin, loaded={"Fed": 1, "OCC": 1, "FinCEN": 1})
    rep, html_out, pdf = email_report.compose(**ev, enforcement_sources=src, edition="Week ending Sun Oct 25")
    names = sorted({m.bank.name for m in rep.enforcement})
    assert names == ["Community Bank 10", "Community Bank 2", "First Prairie Bank"]
    assert rep.counts["ltwc"] == 1 and rep.ltwc[0][0].name == "First Prairie Bank"
    assert "Less than well capitalized" in html_out and "Enforcement orders and capital" in html_out
    assert "checked against Fed" in html_out and rep.subject.startswith("FFIEC Call Report Monitor · Week ending Sun Oct 25 · ")
    assert ">ENFORCEMENT / AML<" in html_out and "BELOW WELL CAP." in html_out
    assert rep.summary.startswith("1 filing is below well capitalized: First Prairie Bank")


def test_old_penalties_dropped_orders_kept():
    from datetime import timedelta
    from ffiec_monitor import enforcement
    from ffiec_monitor.report import _bank
    b = _bank(_meta(1, 1000, "First Prairie Bank"), Q3)
    today = date.today()
    acts = [enforcement.Action("Fed", "First Prairie Bank", "IL", "Civil Money Penalty", today - timedelta(days=2000), "u"),
            enforcement.Action("Fed", "First Prairie Bank", "IL", "Civil Money Penalty", today - timedelta(days=100), "u"),
            enforcement.Action("OCC", "First Prairie Bank", "IL", "Formal Agreement", today - timedelta(days=4000), "u")]
    got = [(m.action.action, m.action.stale) for m in enforcement.match([b], acts)]
    assert got == [("Formal Agreement", True)]                 # penalties dropped; old order flagged to confirm
    aml = enforcement.Action("OCC", "X", "IL", "Consent Order", today, "u", "Bank Secrecy Act/Anti-Money Laundering")
    assert aml.topic == "AML" and acts[2].topic == "Safety & soundness"


def test_fdic_orders_match_by_cert(tmp_path):
    from ffiec_monitor import enforcement
    from ffiec_monitor.report import _bank
    f = tmp_path / "o.csv"
    f.write_text("cert,bank,city,state,action,title,issued,docket,url\n"
                 "9001,Some Other Name Bank,X,IL,Prompt Corrective Action,Supervisory PCA Directive,2026-08-28,FDIC-1,u\n"
                 "9002,Bank B,Y,IL,Cease and Desist / Consent Orders,Consent Order,2025-01-01,FDIC-2,u\n")
    acts, _ = enforcement.load_fdic(f)
    banks = [_bank(_meta(1, 1000, "First Prairie Bank"), Q3), _bank(_meta(3, 1000, "Bank C"), Q3)]
    got = enforcement.match(banks, acts)
    assert [(m.bank.name, m.how, m.action.topic) for m in got] == [("First Prairie Bank", "FDIC cert", "Capital")]


def test_flagged_bank_performance(monkeypatch):
    from ffiec_monitor import enforcement, performance
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    assert performance.prior_quarter_end(Q3) == date(2026, 6, 30)
    assert performance.prior_quarter_end(date(2026, 3, 31)) == date(2025, 12, 31)
    cur = {"tier1_leverage_ratio": 8.4, "roaa": -0.2, "noncurrent_loans_oreo_ratio": 2.6,
           "brokered_reciprocal_deposits_to_total_deposits": 18.0}
    prior = {"tier1_leverage_ratio": 9.1, "roaa": 0.3, "noncurrent_loans_oreo_ratio": 1.9,
             "brokered_reciprocal_deposits_to_total_deposits": 12.0}
    seen = []

    def lookup(rssd, p):
        seen.append((rssd, p))
        return prior
    ev = sample_events(0, 0)
    ev.update(watch_list_changes=[], tier1_events=[], watch_list_snapshot=[], tier2_events=[],
              new_filing_events=[email_report.NewFilingEvent(_meta(1, 500_000, "Thin Bank"), Q3, cur),
                                 email_report.NewFilingEvent(_meta(2, 500_000, "Fine Bank"), Q3,
                                                             {"tier1_leverage_ratio": 11, "total_rbc_ratio": 15})])
    src = enforcement.Sources(loaded={"Fed": 0})
    rep, html, pdf = email_report.compose(**ev, enforcement_sources=src, prior_lookup=lookup)
    assert seen == [("100001", date(2026, 6, 30))]                  # only the flagged (CBLR < 9%) bank
    bp = rep.performance[0]
    assert bp.bank.name == "Thin Bank" and bp.why == ["CBLR bank under 9%"]
    assert bp.read.startswith("Mostly worsening: capital down 0.70 pts, losing money, problem loans 2.60% (up 0.70)")
    assert "How the flagged banks are performing" in html and "Q2 2026" in html
    assert "1 mostly worsening" in rep.summary


def test_enforcement_discussion_and_stale_orders(monkeypatch):
    from datetime import timedelta
    from ffiec_monitor import enforcement
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    today = date.today()
    acts = [enforcement.Action("FDIC", "First Prairie Bank", "IL", "Cease and Desist / Consent Orders",
                               today - timedelta(days=400), "https://orders.example/1", "Consent Order", "9001"),
            enforcement.Action("OCC", "Community Bank 2", "IL", "Cease-and-Desist Order", date(2003, 1, 29), "u")]
    ev = sample_events(0, 0)
    ev.update(watch_list_changes=[], tier1_events=[], watch_list_snapshot=[], tier2_events=[],
              new_filing_events=[email_report.NewFilingEvent(_meta(1, 500_000, "First Prairie Bank"), Q3,
                                                             {"tier1_leverage_ratio": 9.5, "roaa": 0.2}),
                                 email_report.NewFilingEvent(_meta(2, 500_000), Q3, {"tier1_leverage_ratio": 10})])
    src = enforcement.Sources(actions=acts, loaded={"FDIC": 1, "OCC": 1})
    rep, html, pdf = email_report.compose(**ev, enforcement_sources=src,
                                          prior_lookup=lambda r, p: {"tier1_leverage_ratio": 9.9, "roaa": 0.4})
    assert rep.counts["enforcement"] == 1                         # the 2003 OCC order is footnoted, not counted
    note, by = rep.enforcement_notes["100001"]
    assert by == "rules" and note.startswith("First Prairie Bank has been under an FDIC cease and desist / consent orders since")
    assert "Not counted: Community Bank 2" in html and "most likely closed" in html
    assert "special cap" in html and "Enforcement orders and capital" in html


def _pages(pdf: bytes) -> int:
    return len(re.findall(rb"/Type /Page[^s]", pdf))


def test_pdf_is_always_five_pages(monkeypatch):
    from datetime import timedelta
    from ffiec_monitor import analytics, enforcement
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    # quiet day
    ev = sample_events(0, 0)
    ev.update(watch_list_changes=[], tier1_events=[], watch_list_snapshot=[], tier2_events=[])
    assert _pages(email_report.compose(**ev)[2]) == 5
    # very busy day: alerts, 60 refilings, 300 new filings (many flagged), orders, capital problems
    weak = {"tier1_leverage_ratio": 4.6, "total_rbc_ratio": 9.0, "roaa": -0.4, "net_loans_leases_to_deposits": 118,
            "brokered_reciprocal_deposits_to_total_deposits": 30, "core_deposits_to_total_deposits": 55}
    cblr = {"tier1_leverage_ratio": 8.4}
    ev = sample_events(60, 0)
    ev["new_filing_events"] = [email_report.NewFilingEvent(_meta(500 + i, 90_000 + i * 1000, f"Bank {i}"), Q3,
                                                            weak if i % 3 == 0 else cblr if i % 3 == 1 else {"tier1_leverage_ratio": 11})
                               for i in range(300)]
    acts = [enforcement.Action("FDIC", f"Bank {i}", "IL", "Cease and Desist / Consent Orders",
                               date.today() - timedelta(days=200), "u", "Consent Order", str(9500 + i)) for i in range(0, 60, 3)]
    a = analytics.Analytics(run_date=date(2026, 10, 26), period_end=Q3, filed=1500, universe=4343, deadline=date(2026, 10, 30))
    a.days = [(date(2026, 10, 6) + timedelta(days=i), 100 + i, 5) for i in range(21)]
    rep, html, pdf = email_report.compose(**ev, analytics=a, enforcement_sources=enforcement.Sources(actions=acts, loaded={"FDIC": 20}),
                                          prior_lookup=lambda r, p: {"tier1_leverage_ratio": 9.0, "roaa": 0.5})
    assert rep.counts["enforcement"] == 20 and rep.ltwc and rep.cblr and rep.flagged_total > 50
    assert _pages(pdf) == 5


def test_spreadsheet_attachment(monkeypatch):
    import io
    from openpyxl import load_workbook
    from ffiec_monitor.workbook import build_workbook
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    ev = sample_events(3, 0)
    ev["new_filing_events"] = [email_report.NewFilingEvent(_meta(1, 500_000, "Weak Bank"), Q3,
                                                            {"tier1_leverage_ratio": 4.7, "total_rbc_ratio": 9.1}),
                               email_report.NewFilingEvent(_meta(2, 90_000, "Cblr Bank"), Q3, {"tier1_leverage_ratio": 8.5})]
    rep, _, _ = email_report.compose(**ev)
    wb = load_workbook(io.BytesIO(build_workbook(rep, ev["restatement_events"], ev["new_filing_events"])))
    assert wb.sheetnames == ["Summary", "Filings", "Enforcement", "Capital watch", "Flagged performance", "Notes"]
    fs = wb["Filings"]
    head = [c.value for c in fs[4]]
    rows = {r[1]: dict(zip(head, r)) for r in fs.iter_rows(min_row=5, values_only=True)}
    assert rows["Weak Bank"]["Capital status"] == "Below well capitalized" and abs(rows["Weak Bank"]["Leverage ratio"] - 0.047) < 1e-9
    assert rows["Cblr Bank"]["Capital status"] == "CBLR under 9%" and len(rows) == 5
    assert str(wb["Summary"]["B8"].value).startswith('=COUNTIF(Filings!')          # live formulas, not pasted numbers
    assert wb["Summary"]["B8"].offset(0, -1).value == "Below well capitalized"


def test_email_does_not_say_no_orders_when_there_are_some(monkeypatch):
    from ffiec_monitor import enforcement
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    ev = sample_events(0, 0)
    ev.update(watch_list_changes=[], tier1_events=[], watch_list_snapshot=[], tier2_events=[],
              new_filing_events=[email_report.NewFilingEvent(_meta(1, 500_000, "First Prairie Bank"), Q3, {"tier1_leverage_ratio": 9.5})])
    acts = [enforcement.Action("FDIC", "First Prairie Bank", "IL", "Consent Order", date(2025, 3, 1), "u", "", "9001")]
    _, html, _ = email_report.compose(**ev, enforcement_sources=enforcement.Sources(actions=acts, loaded={"FDIC": 1}))
    assert "First Prairie Bank" in html and "No bank is under an active" not in html


def test_ubpr_links():
    from ffiec_monitor.ubpr import ubpr_url
    assert ubpr_url("142179", Q3) == "https://cdr.ffiec.gov/Public/Reports/UbprReport.aspx?rptCycleIds=154&rptid=283&idrssd=142179"
    assert ubpr_url("142179", date(2026, 6, 30)).endswith("rptCycleIds=152&rptid=283&idrssd=142179")
    assert ubpr_url("142179", date(2031, 3, 31), discover=False).endswith("ManageFacsimiles.aspx")
    ev = sample_events(2, 0)
    _, html, _ = email_report.compose(**ev)
    assert "UbprReport.aspx?rptCycleIds=" in html and ">UBPR<" in html
