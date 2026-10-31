"""The workbook's IntraFi ratios, flags and Ratio Guide (ffiec_monitor/filings_workbook.py)."""
import io
from datetime import date

import openpyxl

from ffiec_monitor.filings_workbook import RATIO_COLUMNS, RATIO_GUIDE, build, extra_ratios, intrafi_flags

PERIOD = date(2026, 9, 30)

# A $500M bank ($000) that doesn't report 5597: uninsured is estimated from accounts over $250K.
CONCEPTS = {
    "RCON2170": 500_000, "RCON2200": 400_000, "RCONF051": 180_000, "RCONF052": 200,  # 180M in 200 big accounts
    "RCON0081": 5_000, "RCON0071": 20_000, "RCONB987": 5_000, "RCON1773": 60_000, "RCON1771": 18_000,
    "RCON0416": 30_000,
    "RCON1772": 70_000, "RCON1754": 20_000, "RCOA8274": 45_000, "RCON3123": 3_000,
    "RCONF158": 30_000, "RCONF159": 20_000, "RCON1460": 40_000, "RCONF161": 60_000, "RCOA3792": 50_000,
}


def test_new_ratios_from_call_report_items():
    r = extra_ratios(CONCEPTS, PERIOD)
    # uninsured = 180,000 - 250 x 200 = 130,000 -> 32.5% of 400,000 deposits
    assert round(r["uninsured_deposit_share"], 1) == 32.5
    # liquid = 5+20+5+60+18 = 108,000 (pledged not subtracted) vs 130,000 uninsured
    assert round(r["liquid_to_uninsured"], 1) == 83.1
    # unrealized: AFS -10,000 + HTM -2,000 = 12,000 loss / 45,000 Tier 1
    assert round(r["unrealized_loss_to_tier1"], 1) == 26.7
    # CRE 150,000 / 50,000 total capital; construction 50,000 / 50,000
    assert round(r["cre_to_capital"]) == 300 and round(r["construction_to_capital"]) == 100


def test_flags_and_capital_category():
    r = extra_ratios(CONCEPTS, PERIOD)
    r.update(tier1_leverage_ratio=4.5, tier1_rbc_ratio=9.0, total_rbc_ratio=11.0, cet1_ratio=8.0,
             reciprocal_cap_used=85.0)
    assert intrafi_flags(r) == ["Near/over reciprocal cap", "Below well capitalized", "Liquid assets < uninsured",
                                "Securities losses 25%+ of Tier 1", "CRE concentration"]
    healthy = {"tier1_leverage_ratio": 10.0, "uninsured_deposit_share": 20.0, "liquid_to_uninsured": 200.0}
    assert intrafi_flags(healthy) == []


def test_every_ratios_column_is_explained_in_the_guide():
    guide = {row[0] for row in RATIO_GUIDE}
    assert {h for h, _k, _c in RATIO_COLUMNS} <= guide
    assert all(len(row) == 6 and all(row) for row in RATIO_GUIDE)


def test_workbook_shows_flags_and_guide():
    r = extra_ratios(CONCEPTS, PERIOD)
    r["tier1_leverage_ratio"] = 4.5
    rec = dict(name="TEST BANK", city="X", state="VA", rssd="1", cert="2", period=PERIOD, kind="New filing",
               submitted="", enforcement=False, treasury=False, other_lists=[], call_report_url="https://a",
               ubpr_url="https://b", ratios=r)
    wb = openpyxl.load_workbook(io.BytesIO(build([rec], "Oct 30, 2026", 0)))
    filings = wb["Filings"]
    assert filings.cell(row=1, column=12).value == "Capital category"
    assert filings.cell(row=2, column=12).value == "Adequately capitalized"
    assert "CRE concentration" in filings.cell(row=2, column=13).value
    guide = wb["Ratio Guide"]
    assert guide.cell(row=4, column=4).value == "Normal range"
    assert guide.cell(row=5, column=1).value == "Tier 1 leverage"
