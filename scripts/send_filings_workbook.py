"""Email a three-tab .xlsx of the Call Report filings FFIEC shows as submitted in a window
(Eastern): Summary (counts + every bank on the enforcement or Treasury Desk list, with links),
Filings (one row per filing, Call Report + UBPR links), Ratios (capital, earnings, asset
quality per filing). The only attachment is the workbook -- no PDF, no inline images.

Each filing's facsimile is re-pulled so ratios the daily pipeline doesn't store (ROE,
efficiency, ACL/loans, NPA/assets, Texas ratio) can be computed; if a pull fails, the
stored snapshot ratios are used and the extra columns are left blank. No DB writes.

  python scripts/send_filings_workbook.py START END LABEL
  e.g. "2026-10-28 00:00" "2026-10-30 00:00" "Oct 28-29, 2026"
"""
from __future__ import annotations

import html
import io
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ffiec_monitor import db, email_report, ratios, simple_email
from ffiec_monitor.config import EmailConfig, FFIECConfig, get_duckdb_path
from ffiec_monitor.ffiec_client import FFIECClient
from ffiec_monitor.ubpr import ubpr_url
from ffiec_monitor.ratios import _pick, _ratio_pct, annualization_factor, extract_total_assets

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
FLAG_FILL = PatternFill("solid", fgColor="FDECEC")
LINK_FONT = Font(color="1D5FBF", underline="single")

# (header, ratios-dict key, category). Values are percents, as extract_ratios returns them.
RATIO_COLUMNS = [
    ("Tier 1 leverage", "tier1_leverage_ratio", "Capital"),
    ("CET1", "cet1_ratio", "Capital"),
    ("Tier 1 risk-based", "tier1_rbc_ratio", "Capital"),
    ("Total risk-based", "total_rbc_ratio", "Capital"),
    ("ROA (annualized)", "roaa", "Earnings"),
    ("ROE (annualized)", "roe", "Earnings"),
    ("Net interest margin", "net_interest_margin", "Earnings"),
    ("Efficiency ratio", "efficiency_ratio", "Earnings"),
    ("NPA / assets", "npa_to_assets", "Asset quality"),
    ("Noncurrent loans / loans", "noncurrent_to_loans", "Asset quality"),
    ("Noncurrent loans + OREO / loans + OREO", "noncurrent_loans_oreo_ratio", "Asset quality"),
    ("Net charge-off rate (annualized)", "charge_off_ratio", "Asset quality"),
    ("ACL / loans", "acl_to_loans", "Asset quality"),
    ("Texas ratio", "texas_ratio", "Asset quality"),
    ("Net loans / deposits", "net_loans_leases_to_deposits", "Liquidity"),
    ("Core deposits / deposits", "core_deposits_to_total_deposits", "Liquidity"),
    ("Brokered + reciprocal / deposits", "brokered_reciprocal_deposits_to_total_deposits", "Liquidity"),
    ("Uninsured deposit share", "uninsured_deposit_share", "Liquidity"),
]


def extra_ratios(c: dict[str, Any], period_end) -> dict[str, Optional[float]]:
    """Ratios extract_ratios doesn't compute, from the same facsimile concepts."""
    factor = annualization_factor(period_end)
    net_income = _pick(c, "RIAD4340")
    equity = _pick(c, "RCFD3210", "RCON3210")
    nii = _pick(c, "RIAD4074")
    nonint_income = _pick(c, "RIAD4079")
    nonint_expense = _pick(c, "RIAD4093")
    gross_loans = _pick(c, "RCFDB528", "RCONB528")
    acl = _pick(c, "RCFD3123", "RCON3123")
    nonaccrual = _pick(c, "RCFD1403", "RCON1403") or 0.0
    past_due_90 = _pick(c, "RCFD1407", "RCON1407") or 0.0
    oreo = _pick(c, "RCFD2150", "RCON2150") or 0.0
    goodwill = _pick(c, "RCFD3163", "RCON3163") or 0.0
    intangibles = _pick(c, "RCFD0426", "RCON0426") or 0.0
    total_assets = extract_total_assets(c)
    out: dict[str, Optional[float]] = {"total_assets": total_assets}
    out["roe"] = _ratio_pct(net_income * factor, equity) if net_income is not None else None
    revenue = (nii or 0.0) + (nonint_income or 0.0)
    out["efficiency_ratio"] = _ratio_pct(nonint_expense, revenue) if nonint_expense is not None and revenue else None
    out["acl_to_loans"] = _ratio_pct(acl, gross_loans)
    out["noncurrent_to_loans"] = _ratio_pct(nonaccrual + past_due_90, gross_loans) if gross_loans else None
    out["npa_to_assets"] = _ratio_pct(nonaccrual + past_due_90 + oreo, total_assets) if total_assets else None
    tangible_plus_acl = (equity - goodwill - intangibles + (acl or 0.0)) if equity is not None else None
    out["texas_ratio"] = _ratio_pct(nonaccrual + past_due_90 + oreo, tangible_plus_acl) if tangible_plus_acl else None
    return out


def _header(ws, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


def _link(cell, url: str, text: str) -> None:
    cell.value, cell.hyperlink, cell.font = text, url, LINK_FONT


def _pct(v) -> Optional[float]:
    return v / 100 if isinstance(v, (int, float)) else None


def build(records: list[dict], label: str, failed: int) -> bytes:
    wb = Workbook()
    flagged = [r for r in records if r["enforcement"] or r["treasury"]]

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Call Report filings {label}"
    ws["A1"].font = Font(bold=True, size=14)
    stats = [
        ("Total filings", len(records)),
        ("New filings", sum(1 for r in records if r["kind"] == "New filing")),
        ("Refilings", sum(1 for r in records if r["kind"] == "Refiling")),
        ("On the enforcement list", sum(1 for r in records if r["enforcement"])),
        ("On the Treasury Desk list", sum(1 for r in records if r["treasury"])),
    ]
    for i, (k, v) in enumerate(stats, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    note = ("Enforcement list = active orders in the monitor's enforcement_orders table plus the FDIC "
            "enforcement-decisions export (data/fdic_active_orders.csv), matched by RSSD / FDIC cert. "
            "Treasury Desk list = data/treasury-desk.csv, matched by bank name + state.")
    if failed:
        note += f" {failed} facsimile pull(s) failed; those rows use stored ratios and leave the extra ratios blank."
    ws.cell(row=9, column=1, value=note).alignment = Alignment(wrap_text=True)
    ws.merge_cells("A9:H9")
    ws.row_dimensions[9].height = 48
    ws.cell(row=11, column=1, value="Banks on the enforcement or Treasury Desk list").font = Font(bold=True, size=12)
    _header(ws, 12, ["Institution", "City", "State", "Type", "Enforcement", "Treasury Desk", "Call Report", "UBPR"])
    ws.freeze_panes = None
    for i, r in enumerate(flagged, 13):
        vals = [r["name"], r["city"], r["state"], r["kind"], "Yes" if r["enforcement"] else "",
                "Yes" if r["treasury"] else ""]
        for col, v in enumerate(vals, 1):
            ws.cell(row=i, column=col, value=v)
        _link(ws.cell(row=i, column=7), r["call_report_url"], "Call Report")
        _link(ws.cell(row=i, column=8), r["ubpr_url"], "UBPR")
    if not flagged:
        ws.cell(row=13, column=1, value="None")
    for col, w in zip("ABCDEFGH", (44, 18, 7, 12, 13, 14, 13, 9)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Filings")
    heads = ["Institution", "City", "State", "RSSD", "FDIC cert", "Period", "Type", "Submitted (ET)",
             "Enforcement", "Treasury Desk", "Other lists", "Call Report", "UBPR"]
    _header(ws, 1, heads)
    for i, r in enumerate(records, 2):
        vals = [r["name"], r["city"], r["state"], r["rssd"], r["cert"], r["period"], r["kind"], r["submitted"],
                "Yes" if r["enforcement"] else "", "Yes" if r["treasury"] else "", ", ".join(r["other_lists"])]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=col, value=v)
            if r["enforcement"] or r["treasury"]:
                cell.fill = FLAG_FILL
        ws.cell(row=i, column=6).number_format = "yyyy-mm-dd"
        _link(ws.cell(row=i, column=12), r["call_report_url"], "Call Report")
        _link(ws.cell(row=i, column=13), r["ubpr_url"], "UBPR")
    for col, w in zip("ABCDEFGHIJKLM", (44, 18, 7, 10, 10, 11, 11, 17, 12, 13, 18, 12, 8)):
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A1:M{len(records) + 1}"

    ws = wb.create_sheet("Ratios")
    heads = ["Institution", "State", "RSSD", "Period", "Total assets ($000)"] + [h for h, _, _ in RATIO_COLUMNS]
    ws.cell(row=1, column=1, value="")
    for col, (_, _, cat) in enumerate(RATIO_COLUMNS, 6):
        ws.cell(row=1, column=col, value=cat).font = Font(bold=True, color="555555")
    _header(ws, 2, heads)
    ws.row_dimensions[2].height = 45
    for i, r in enumerate(records, 3):
        for col, v in enumerate([r["name"], r["state"], r["rssd"], r["period"], r["ratios"].get("total_assets")], 1):
            ws.cell(row=i, column=col, value=v)
        ws.cell(row=i, column=4).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=5).number_format = "#,##0"
        for col, (_, key, _) in enumerate(RATIO_COLUMNS, 6):
            cell = ws.cell(row=i, column=col, value=_pct(r["ratios"].get(key)))
            cell.number_format = "0.00%"
    ws.column_dimensions["A"].width = 44
    for col in range(2, len(heads) + 1):
        ws.column_dimensions[ws.cell(row=2, column=col).column_letter].width = 13
    ws.auto_filter.ref = f"A2:{ws.cell(row=2, column=len(heads)).column_letter}{len(records) + 2}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main(start: str, end: str, label: str) -> None:
    conn = db.connect(get_duckdb_path())
    filings = simple_email.filings_submitted_between(conn, datetime.fromisoformat(start), datetime.fromisoformat(end))
    checker = simple_email.ListChecker(conn)
    client = FFIECClient(FFIECConfig.from_env())
    records, failed = [], 0
    for f in filings:
        meta = db.get_institution_meta(conn, f.id_rssd) or db.InstitutionMeta(
            id_rssd=f.id_rssd, name=f"RSSD {f.id_rssd}", city="", state="", fdic_cert="",
            total_assets=None, as_of_period=f.period_end)
        try:
            fac = client.retrieve_facsimile(f.id_rssd, f.period_end)
            r = ratios.extract_ratios(fac, f.period_end)
            r.update(extra_ratios(fac.get("concepts", fac), f.period_end))
        except Exception as exc:
            failed += 1
            print(f"facsimile pull failed for {f.id_rssd}: {exc}")
            row = conn.execute("SELECT ratios FROM call_report_snapshots WHERE id_rssd = ? AND period_end = ?",
                               [f.id_rssd, f.period_end]).fetchone()
            r = json.loads(row[0]) if row and row[0] else {}
            r["total_assets"] = meta.total_assets
        lists = checker.lists_for(meta)
        records.append({
            "name": meta.name.strip(), "city": meta.city.strip(), "state": meta.state.strip(),
            "rssd": f.id_rssd, "cert": (meta.fdic_cert or "").strip(), "period": f.period_end,
            "kind": "Refiling" if f.refiling else "New filing",
            "submitted": f"{f.submitted:%Y-%m-%d %I:%M %p}" if f.submitted else "",
            "enforcement": "FDIC order" in lists, "treasury": "Treasury Desk" in lists,
            "other_lists": [x for x in lists if x not in ("FDIC order", "Treasury Desk")],
            "call_report_url": email_report.cdr_report_url(f.id_rssd, f.period_end),
            "ubpr_url": ubpr_url(f.id_rssd, f.period_end),
            "ratios": r,
        })
    records.sort(key=lambda x: (not (x["enforcement"] or x["treasury"]), x["name"].upper()))

    enf = sum(1 for x in records if x["enforcement"])
    tsy = sum(1 for x in records if x["treasury"])
    subject = f"Call Report workbook {label}: {len(records)} filings, {enf} enforcement, {tsy} Treasury Desk"
    text = (f"Call Report filings {label}\n\nTotal filings: {len(records)}\nOn the enforcement list: {enf}\n"
            f"On the Treasury Desk list: {tsy}\n\nThe attached workbook has three tabs: Summary, Filings "
            f"(Call Report + UBPR links) and Ratios (capital, earnings, asset quality).")
    if failed:
        text += f"\n\n{failed} facsimile pull(s) failed; those rows use stored ratios."
    print(subject)
    print(text)
    data = build(records, label, failed)
    fname = f"Call-Report-Filings-{label.replace(' ', '-').replace(',', '')}.xlsx"
    email_report.send_email(EmailConfig.from_env(), subject, text,
                            html_body="<div style='font-family:Arial,sans-serif;font-size:13px;'>"
                                      + html.escape(text).replace("\n", "<br>") + "</div>",
                            attachments=[(fname, data, "vnd.openxmlformats-officedocument.spreadsheetml.sheet")])
    print(f"Sent {fname} ({len(data):,} bytes)")


if __name__ == "__main__":
    main(*sys.argv[1:4])
